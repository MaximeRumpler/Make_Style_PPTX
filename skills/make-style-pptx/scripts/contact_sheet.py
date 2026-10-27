"""Contact sheet of the shapes library: every container shape at three sizes.

Each shape shows its safe text box (dashed) filled with sample text cut, at a word
boundary, to what that box holds. A kind with a minimum size (the drop) is shown at
its smallest usable size for each font size, never below. Colours are theme colours only (the accent ramp of the
Make_Clean theme contract, dark to light), fonts are theme fonts.

    python contact_sheet.py                      # docs/shapes-contact-sheet.pptx + .png
    python contact_sheet.py --out sheet.pptx --png sheet.png

The PNG needs LibreOffice (soffice) and pdftoppm; without them only the PPTX is written.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import _shapes as S  # noqa: E402

ICONS = HERE.parent / "assets" / "icons"
REPO = HERE.parents[2]

# Neutral values of the Make_Clean theme contract (dark-to-light accent ramp).
THEME = {"dk1": "111111", "lt1": "FFFFFF", "dk2": "404040", "lt2": "F2F2F2",
         "accent1": "1F1F1F", "accent2": "595959", "accent3": "8C8C8C", "accent4": "B3B3B3",
         "accent5": "D0D0D0", "accent6": "E6E6E6", "hlink": "404040", "folHlink": "8C8C8C"}
FONT = "Arial"

# (label, kind, params, aspect for box shapes or None for the natural one)
SHEET = [
    ("rectangle", "rectangle", {}, 1.6),
    ("rounded", "rounded", {}, 1.5),
    ("pill", "pill", {}, 3.2),
    ("ellipse", "ellipse", {}, 1.45),
    ("circle", "circle", {}, 1.0),
    ("drop", "drop", {}, None),
    ("drop, tilt 6", "drop", {"tilt": 6}, None),
    ("cloud", "cloud", {}, None),
    ("cloud, 4 lobes, flipped", "cloud", {"lobes": 4, "flip": True}, None),
    ("blob, seed 1", "blob", {"seed": 1}, None),
    ("blob, seed 4, wobble 0.14", "blob", {"seed": 4, "wobble": 0.14}, 1.35),
    ("pebble, seed 2", "pebble", {"seed": 2}, 2.0),
    ("pebble, seed 5, 2.5:1", "pebble", {"seed": 5}, 2.5),
    ("ticket", "ticket", {}, 2.2),
    ("ticket, corner notches", "ticket", {"style": "corners"}, 1.9),
    ("speech bubble", "bubble", {}, 1.5),
    ("speech bubble, tail left", "bubble", {"side": "left", "pos": 0.65, "lean": 0.5}, 1.75),
    ("speech bubble, tail right", "bubble", {"side": "right", "pos": 0.3, "lean": -0.4}, 1.75),
    ("torn paper", "torn", {}, 1.25),
    ("torn paper, top and bottom", "torn", {"edges": "top+bottom", "seed": 3}, 1.9),
    ("hexagon", "polygon", {"sides": 6}, None),
    ("pentagon", "polygon", {"sides": 5}, None),
    ("hand-drawn", "handdrawn", {}, 1.6),
    ("hand-drawn, seed 7", "handdrawn", {"seed": 7}, 1.15),
    ("hand-drawn diamond", "handdrawn", {"seed": 3, "form": "diamond", "strokes": 1}, 1.0),
    ("icon: shield (Phosphor)", "icon", {"svg": str(ICONS / "phosphor" / "shield-fill.svg")}, None),
    ("icon: seal (Phosphor)", "icon", {"svg": str(ICONS / "phosphor" / "seal-fill.svg")}, None),
    ("icon: chat (Phosphor)", "icon", {"svg": str(ICONS / "phosphor" / "chat-circle-fill.svg")}, None),
    ("icon: bookmark (Tabler)", "icon", {"svg": str(ICONS / "tabler" / "bookmark.svg")}, None),
    ("icon: shield (Tabler)", "icon", {"svg": str(ICONS / "tabler" / "shield.svg")}, None),
]

SAMPLE = ("Orders held steady through a grey quarter while every lost client got a call and costs "
          "stayed flat without a single cut, which left room to plan the spring launch with care.")

# size steps: share of the large cell, fill, text colour, font size
SIZES = [(1.0, MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.BACKGROUND_1, 14),
         (0.72, MSO_THEME_COLOR.ACCENT_3, MSO_THEME_COLOR.TEXT_1, 11),
         (0.5, MSO_THEME_COLOR.ACCENT_5, MSO_THEME_COLOR.TEXT_1, 9)]
DASH = {MSO_THEME_COLOR.ACCENT_1: MSO_THEME_COLOR.ACCENT_4,
        MSO_THEME_COLOR.ACCENT_3: MSO_THEME_COLOR.BACKGROUND_1,
        MSO_THEME_COLOR.ACCENT_5: MSO_THEME_COLOR.ACCENT_2}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
COLS = 2
CAPTION_W = Inches(0.95)
CELL_W = Inches(6.35)
LABEL_H, CAPTION_H, ROW_GAP = Inches(0.3), Inches(0.42), Inches(0.05)
TOP, BOTTOM = Inches(0.75), Inches(7.45)
MAX_W, MAX_H = Inches(2.6), Inches(1.45)


def _set_theme(prs) -> None:
    master = prs.slide_masters[0].part
    theme = master.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    ns = {"a": S.A_NS}
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", "Neutral ramp")
    for slot, hexval in THEME.items():
        el = scheme.find(f"a:{slot}", ns)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{S.A_NS}}}srgbClr", val=hexval)
    for which in ("majorFont", "minorFont"):
        root.find(f".//a:{which}/a:latin", ns).set("typeface", FONT)
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _box_for(kind, params, aspect, scale, size=S.REFERENCE_PT):
    a = aspect if aspect is not None else S.natural_aspect(kind, **params)
    w, h = MAX_W, MAX_W / a
    if h > MAX_H:
        h, w = MAX_H, MAX_H * a
    w, h = w * scale, h * scale
    ms = S.min_size(kind, size)
    if ms is not None:  # never below the smallest usable size for this font size
        k = max(ms[0] / w, ms[1] / h, 1.0)
        w, h = w * k, h * k
    return int(w), int(h)


def _text(slide, x, y, w, h, text, size, colour, *, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, dash=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.theme_color = colour
    if dash is not None:
        tb.line.color.theme_color = dash
        tb.line.width = Pt(0.75)
        ln = tb.line._get_or_add_ln()
        etree.SubElement(ln, f"{{{S.A_NS}}}prstDash", val="dash")
    return tb


def _row_height(cells) -> int:
    tallest = max(_box_for(kind, params, aspect, scale, size)[1]
                  for _, kind, params, aspect in cells for scale, _, _, size in SIZES)
    return LABEL_H + tallest + CAPTION_H


def _pages():
    """Rows of COLS cells, flowed onto slides; a row is as tall as its tallest shape."""
    rows = [SHEET[i:i + COLS] for i in range(0, len(SHEET), COLS)]
    pages, cur, y = [], [], TOP
    for row in rows:
        rh = _row_height(row)
        if cur and y + rh > BOTTOM:
            pages.append(cur)
            cur, y = [], TOP
        cur.append((y, row, rh))
        y += rh + ROW_GAP
    pages.append(cur)
    return pages


def build(out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _set_theme(prs)
    blank = prs.slide_layouts[6]
    pages = _pages()
    margin_x = Inches(0.33)
    for page_no, page in enumerate(pages, 1):
        slide = prs.slides.add_slide(blank)
        _text(slide, margin_x, Inches(0.25), Inches(9), Inches(0.4), "Shapes library: container silhouettes",
              18, MSO_THEME_COLOR.TEXT_1, bold=True)
        _text(slide, Inches(8.5), Inches(0.3), Inches(4.5), Inches(0.35),
              f"Dashed: safe text box, sample text cut at a word to what it holds · {page_no}/{len(pages)}", 9,
              MSO_THEME_COLOR.ACCENT_3, align=PP_ALIGN.RIGHT)
        for cy, row, rh in page:
            for i, (label, kind, params, aspect) in enumerate(row):
                cx = margin_x + i * CELL_W
                _text(slide, cx, cy, CELL_W, LABEL_H, label, 11, MSO_THEME_COLOR.TEXT_1, bold=True)
                x = cx
                base = cy + rh - CAPTION_H  # shapes sit on a common baseline
                for scale, fill, ink, size in SIZES:
                    w, h = _box_for(kind, params, aspect, scale, size)
                    y = base - h
                    shp, ins = S.add_shape(slide, kind, (x, y, w, h), name=f"{kind} {scale}", **params)
                    shp.fill.solid()
                    shp.fill.fore_color.theme_color = fill
                    if kind == "handdrawn":
                        shp.line.color.theme_color = MSO_THEME_COLOR.TEXT_1
                        shp.line.width = Pt(1.25)
                    tx, ty, tw, th = ins.rect(x, y, w, h)
                    words = S.word_limit(kind, (w, h), size, **params)
                    text = S.fit_words(kind, (w, h), SAMPLE, size, **params)
                    _text(slide, tx, ty, tw, th, text, size, ink, align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE, dash=DASH[fill])
                    _text(slide, x, base + Inches(0.06), max(w, CAPTION_W), Inches(0.3),
                          f"{w / S.EMU_PER_INCH:.2f} x {h / S.EMU_PER_INCH:.2f} in\n"
                          + (f"{S.char_limit(tw, th, size)} chars, {words} word{'' if words == 1 else 's'} at {size} pt"
                           if text else f"too small for this text at {size} pt"),
                          7, MSO_THEME_COLOR.ACCENT_3)
                    x += max(w, CAPTION_W) + Inches(0.25)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def render(pptx: Path, png: Path, dpi: int = 72) -> Path | None:
    """Render every slide and stack them into one PNG. Returns None when LibreOffice
    or pdftoppm is missing."""
    office, pdftoppm = soffice(), shutil.which("pdftoppm")
    if not office or not pdftoppm:
        return None
    from PIL import Image

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        subprocess.run([office, "--headless", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "pdf",
                        "--outdir", str(tmp), str(pptx)], check=True, capture_output=True, timeout=240)
        pdf = tmp / (pptx.stem + ".pdf")
        subprocess.run([pdftoppm, "-png", "-r", str(dpi), str(pdf), str(tmp / "page")], check=True,
                       capture_output=True, timeout=240)
        pages = sorted(tmp.glob("page*.png"))
        imgs = [Image.open(p).convert("RGB") for p in pages]
        gap = 12
        sheet = Image.new("RGB", (max(i.width for i in imgs), sum(i.height for i in imgs) + gap * (len(imgs) - 1)),
                          (200, 200, 200))
        y = 0
        for im in imgs:
            sheet.paste(im, (0, y))
            y += im.height + gap
        png.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(png, optimize=True)
    return png


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=REPO / "docs" / "shapes-contact-sheet.pptx")
    ap.add_argument("--png", type=Path, default=REPO / "docs" / "shapes-contact-sheet.png")
    ap.add_argument("--dpi", type=int, default=72)
    args = ap.parse_args(argv)
    build(args.out)
    print(f"wrote {args.out}")
    if render(args.out, args.png, args.dpi):
        print(f"wrote {args.png}")
    else:
        print("LibreOffice or pdftoppm not found: PNG skipped")
    return 0


if __name__ == "__main__":
    sys.exit(main())
