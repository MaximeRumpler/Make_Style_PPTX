"""Container silhouettes as native, editable PowerPoint geometry.

Every shape is drawn as DrawingML custom geometry (``a:custGeom``) made of lines and
smooth cubic Bezier curves, sized for the box it fills. It stays a normal shape: it
takes theme colours, can be edited point by point in PowerPoint or LibreOffice, and
holds text. For each shape the module also computes the **safe text box**: the
largest axis-aligned rectangle inside the silhouette, returned as insets from the
box, so text never spills out of a drop, a cloud or a speech bubble.

Entry point for the restyler::

    geom, insets = make_shape("drop", (x, y, w, h), tilt=6)

``geom`` is an ``a:custGeom`` lxml element (its text rectangle is already set to
the safe box), ``insets`` the safe text box as EMU insets from the shape box.
``add_shape`` puts one on a python-pptx slide. ``fits`` says whether a shape can be used at a box (aspect bounds, the kind's
``min_size`` at the font size, and, given the text, whole words only);
``aspect_bounds`` gives the proportions it accepts without looking distorted;
``char_limit`` turns a text box into a character limit with the style contract
formula, ``word_limit`` into a whole-word limit, and ``fit_words`` / ``cut_words``
cut text at word boundaries, never mid-word.

Coordinates are EMU (914,400 per inch), y pointing down. Only the standard library,
lxml and Pillow are used; output is deterministic (same parameters, same bytes).
"""
from __future__ import annotations

import math
import random
from pathlib import Path as _FsPath
from typing import Callable, Dict, List, NamedTuple, Optional, Sequence, Tuple

from lxml import etree

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
EMU_PER_INCH = 914400
EMU_PER_PT = 12700

Seg = tuple  # ("M", x, y) | ("L", x, y) | ("C", x1, y1, x2, y2, x, y) | ("Z",)
Sub = List[Seg]


class ShapeError(ValueError):
    """Unknown shape, bad parameter, or a box the shape cannot fill without distortion."""


class Insets(NamedTuple):
    """Safe text box as insets (EMU) from the left, top, right and bottom of the shape box."""

    left: int
    top: int
    right: int
    bottom: int

    def rect(self, x: int, y: int, w: int, h: int) -> Tuple[int, int, int, int]:
        """The text box (x, y, w, h) for a shape placed at (x, y, w, h)."""
        return x + self.left, y + self.top, w - self.left - self.right, h - self.top - self.bottom

    def fractions(self, w: int, h: int) -> Tuple[float, float, float, float]:
        return self.left / w, self.top / h, self.right / w, self.bottom / h


class ShapeResult(NamedTuple):
    geom: etree._Element  # a:custGeom
    insets: Insets


class Path(NamedTuple):
    """One a:path: subpaths plus how it is painted. ``fill=False`` paths are outlines only
    (a sketch stroke) and do not count for the text box."""

    subs: List[Sub]
    fill: bool = True
    stroke: bool = True


Geometry = List[Path]


# ------------------------------------------------------------------ pen
class Pen:
    """Small path builder. Angles are in radians, y down: 0 = right, pi/2 = bottom."""

    def __init__(self) -> None:
        self.subs: List[Sub] = []
        self.x = self.y = 0.0

    def move(self, x: float, y: float) -> "Pen":
        self.subs.append([("M", x, y)])
        self.x, self.y = x, y
        return self

    def line(self, x: float, y: float) -> "Pen":
        self.subs[-1].append(("L", x, y))
        self.x, self.y = x, y
        return self

    def curve(self, x1, y1, x2, y2, x, y) -> "Pen":
        self.subs[-1].append(("C", x1, y1, x2, y2, x, y))
        self.x, self.y = x, y
        return self

    def arc(self, cx: float, cy: float, rx: float, ry: float, a0: float, a1: float) -> "Pen":
        """Elliptical arc from angle a0 to a1 (the pen must stand on the a0 point)."""
        n = max(1, int(math.ceil(abs(a1 - a0) / (math.pi / 2) - 1e-9)))
        step = (a1 - a0) / n
        k = 4.0 / 3.0 * math.tan(step / 4.0)
        t = a0
        for _ in range(n):
            c0, s0 = math.cos(t), math.sin(t)
            c1, s1 = math.cos(t + step), math.sin(t + step)
            self.curve(cx + rx * (c0 - k * s0), cy + ry * (s0 + k * c0),
                       cx + rx * (c1 + k * s1), cy + ry * (s1 - k * c1),
                       cx + rx * c1, cy + ry * s1)
            t += step
        return self

    def close(self) -> "Pen":
        self.subs[-1].append(("Z",))
        return self


def _smooth_closed(points: Sequence[Tuple[float, float]], tension: float = 1.0) -> Sub:
    """Closed Catmull-Rom spline through the points, as cubic Beziers (C1-continuous)."""
    n = len(points)
    sub: Sub = [("M",) + tuple(points[0])]
    for i in range(n):
        p0, p1, p2, p3 = points[i - 1], points[i], points[(i + 1) % n], points[(i + 2) % n]
        c1 = (p1[0] + (p2[0] - p0[0]) * tension / 6, p1[1] + (p2[1] - p0[1]) * tension / 6)
        c2 = (p2[0] - (p3[0] - p1[0]) * tension / 6, p2[1] - (p3[1] - p1[1]) * tension / 6)
        sub.append(("C", c1[0], c1[1], c2[0], c2[1], p2[0], p2[1]))
    sub.append(("Z",))
    return sub


def _map(geom: Geometry, f: Callable[[float, float], Tuple[float, float]]) -> Geometry:
    out = []
    for p in geom:
        subs = []
        for sub in p.subs:
            ns = []
            for s in sub:
                if s[0] == "Z":
                    ns.append(s)
                    continue
                pts = []
                for k in range(1, len(s), 2):
                    pts.extend(f(s[k], s[k + 1]))
                ns.append((s[0],) + tuple(pts))
            subs.append(ns)
        out.append(Path(subs, p.fill, p.stroke))
    return out


def _cubic_extrema(a: float, b: float, c: float, d: float) -> List[float]:
    """Parameters t in (0, 1) where the cubic Bezier coordinate has a local extremum."""
    qa, qb, qc = -a + 3 * b - 3 * c + d, 2 * (a - 2 * b + c), b - a  # derivative / 3
    ts = []
    if abs(qa) < 1e-12:
        if abs(qb) > 1e-12:
            ts.append(-qc / qb)
    else:
        disc = qb * qb - 4 * qa * qc
        if disc >= 0:
            r = math.sqrt(disc)
            ts += [(-qb + r) / (2 * qa), (-qb - r) / (2 * qa)]
    return [t for t in ts if 0 < t < 1]


def _bez(a, b, c, d, t):
    u = 1 - t
    return u * u * u * a + 3 * u * u * t * b + 3 * u * t * t * c + t * t * t * d


def bbox(geom: Geometry) -> Tuple[float, float, float, float]:
    """Exact bounding box (x0, y0, x1, y1) of the geometry, curves included."""
    xs, ys = [], []
    for p in geom:
        for sub in p.subs:
            cx = cy = 0.0
            for s in sub:
                if s[0] in ("M", "L"):
                    cx, cy = s[1], s[2]
                    xs.append(cx)
                    ys.append(cy)
                elif s[0] == "C":
                    for t in _cubic_extrema(cx, s[1], s[3], s[5]):
                        xs.append(_bez(cx, s[1], s[3], s[5], t))
                    for t in _cubic_extrema(cy, s[2], s[4], s[6]):
                        ys.append(_bez(cy, s[2], s[4], s[6], t))
                    cx, cy = s[5], s[6]
                    xs.append(cx)
                    ys.append(cy)
    return min(xs), min(ys), max(xs), max(ys)


def _fit(geom: Geometry, w: float, h: float) -> Geometry:
    """Stretch a design-space geometry so its bounding box fills (0, 0, w, h)."""
    x0, y0, x1, y1 = bbox(geom)
    sx, sy = w / (x1 - x0), h / (y1 - y0)
    return _map(geom, lambda x, y: ((x - x0) * sx, (y - y0) * sy))


def _rotate(geom: Geometry, degrees: float) -> Geometry:
    a = math.radians(degrees)
    c, s = math.cos(a), math.sin(a)
    return _map(geom, lambda x, y: (x * c - y * s, x * s + y * c))


def flatten(sub: Sub, steps: int = 24) -> List[Tuple[float, float]]:
    """Polygon approximation of one subpath (for rasterising and tests)."""
    pts: List[Tuple[float, float]] = []
    cx = cy = 0.0
    for s in sub:
        if s[0] in ("M", "L"):
            cx, cy = s[1], s[2]
            pts.append((cx, cy))
        elif s[0] == "C":
            for i in range(1, steps + 1):
                t = i / steps
                pts.append((_bez(cx, s[1], s[3], s[5], t), _bez(cy, s[2], s[4], s[6], t)))
            cx, cy = s[5], s[6]
    return pts


# ------------------------------------------------------------ the shapes
# Two families. "Box" builders draw directly at the box size (w, h) so corners stay
# round and notches stay circular at any proportion. "Design" builders draw in their
# own units at their natural proportion; the result is stretched to the box, and the
# aspect bounds keep that stretch small.

def _rect_geo(w, h, **_):
    return [Path([Pen().move(0, 0).line(w, 0).line(w, h).line(0, h).close().subs[0]])]


def _rounded_rect(pen: Pen, x, y, w, h, r):
    r = max(0.0, min(r, w / 2, h / 2))
    if r <= 0:
        return pen.move(x, y).line(x + w, y).line(x + w, y + h).line(x, y + h).close()
    pen.move(x + r, y).line(x + w - r, y)
    pen.arc(x + w - r, y + r, r, r, -math.pi / 2, 0).line(x + w, y + h - r)
    pen.arc(x + w - r, y + h - r, r, r, 0, math.pi / 2).line(x + r, y + h)
    pen.arc(x + r, y + h - r, r, r, math.pi / 2, math.pi).line(x, y + r)
    pen.arc(x + r, y + r, r, r, math.pi, 1.5 * math.pi)
    return pen.close()


def _radius(w, h, radius, radius_emu):
    return radius_emu if radius_emu is not None else radius * min(w, h)


def _rounded_geo(w, h, radius=0.18, radius_emu=None, **_):
    return [Path(_rounded_rect(Pen(), 0, 0, w, h, _radius(w, h, radius, radius_emu)).subs)]


def _pill_geo(w, h, **_):
    return [Path(_rounded_rect(Pen(), 0, 0, w, h, min(w, h) / 2).subs)]


def _ellipse_geo(w, h, **_):
    pen = Pen().move(w, h / 2).arc(w / 2, h / 2, w / 2, h / 2, 0, 2 * math.pi).close()
    return [Path(pen.subs)]


def _circle_geo(w, h, **_):
    r = min(w, h) / 2
    pen = Pen().move(w / 2 + r, h / 2).arc(w / 2, h / 2, r, r, 0, 2 * math.pi).close()
    return [Path(pen.subs)]


DROP_MAX_TILT = 8.0  # beyond this a leaning drop reads as a flame


def _drop_geo(tilt=0.0, body=0.63, tip=0.5, **_):
    """Teardrop, tip up. ``body``: diameter of the round body over the total height
    (0.55 to 0.75); ``tip``: 0 for a blunt tip, 1 for a needle-sharp one; ``tilt``:
    degrees, positive leans the tip to the right, clamped to +/- DROP_MAX_TILT.

    The sides leave the tip, swell outwards (convex, as surface tension would draw
    them) and meet the body at its widest point with a vertical tangent, so the outline
    is smooth everywhere except at the tip itself.
    """
    if not 0.55 <= body <= 0.75:
        raise ShapeError("drop: body must be between 0.55 and 0.75 of the height")
    if not 0 <= tip <= 1:
        raise ShapeError("drop: tip must be between 0 and 1")
    r = body / 2
    cy = 1.0 - r  # centre of the body; the widest point is at this height
    beta = math.radians(46 - 24 * tip)  # angle of the side at the tip, from the vertical
    d1 = cy * (0.31 - 0.19 * tip)  # handle at the tip
    d2 = cy * 0.48  # handle above the widest point (fullness of the shoulders)
    pen = Pen().move(0, 0)
    pen.curve(d1 * math.sin(beta), d1 * math.cos(beta), r, cy - d2, r, cy)
    pen.arc(0, cy, r, r, 0, math.pi)
    pen.curve(-r, cy - d2, -d1 * math.sin(beta), d1 * math.cos(beta), 0, 0)
    geom = [Path(pen.close().subs)]
    tilt = max(-DROP_MAX_TILT, min(DROP_MAX_TILT, float(tilt)))
    return _rotate(geom, tilt) if tilt else geom


def _circle_hits(c1, c2, upper=True):
    (x1, y1, r1), (x2, y2, r2) = c1, c2
    dx, dy = x2 - x1, y2 - y1
    d = math.hypot(dx, dy)
    if d >= r1 + r2 or d <= abs(r1 - r2):
        raise ShapeError("cloud: lobes do not overlap")
    a = (r1 * r1 - r2 * r2 + d * d) / (2 * d)
    hh = math.sqrt(max(0.0, r1 * r1 - a * a))
    mx, my = x1 + a * dx / d, y1 + a * dy / d
    p = (mx + hh * dy / d, my - hh * dx / d)
    q = (mx - hh * dy / d, my + hh * dx / d)
    return min(p, q, key=lambda t: t[1]) if upper else max(p, q, key=lambda t: t[1])


_CLOUDS = {
    # lobes (cx, cy, r) from right to left; the outer two rest on the flat base
    3: [(0.765, 0.42, 0.20), (0.44, 0.33, 0.255), (0.19, 0.48, 0.14)],
    4: [(0.835, 0.47, 0.15), (0.635, 0.335, 0.215), (0.375, 0.315, 0.205), (0.17, 0.465, 0.155)],
}


def _cloud_geo(lobes=3, flip=False, **_):
    """Flat-bottomed cloud: overlapping round lobes, the outer two resting on the base,
    the largest one off-centre so the silhouette does not look stamped."""
    if lobes not in _CLOUDS:
        raise ShapeError("cloud: lobes must be 3 or 4")
    cs = _CLOUDS[lobes]
    first, last = cs[0], cs[-1]
    base = first[1] + first[2]
    hits = [_circle_hits(cs[i], cs[i + 1]) for i in range(len(cs) - 1)]
    pen = Pen().move(last[0], base).line(first[0], base)

    def ang(c, p):
        return math.atan2(p[1] - c[1], p[0] - c[0])

    for i, c in enumerate(cs):
        start = math.pi / 2 if i == 0 else ang(c, hits[i - 1])
        end = ang(c, hits[i]) if i < len(hits) else math.pi / 2
        while end >= start:
            end -= 2 * math.pi
        pen.arc(c[0], c[1], c[2], c[2], start, end)
    geom = [Path(pen.close().subs)]
    return _map(geom, lambda x, y: (-x, y)) if flip else geom


def _blob_geo(seed=0, wobble=0.10, **_):
    """Organic blob: a circle whose radius varies with a few slow harmonics (random
    phases from ``seed``), traced by a smooth closed spline. Low harmonics only, so it
    reads as a pebble or a soft stain, never as a star."""
    if not 0 <= wobble <= 0.2:
        raise ShapeError("blob: wobble must be between 0 and 0.2")
    rnd = random.Random(int(seed))
    harm = [(2, 1.0), (3, 0.55), (4, 0.25)]
    phases = [rnd.random() * 2 * math.pi for _ in harm]
    amps = [a * (0.7 + 0.6 * rnd.random()) for _, a in harm]
    norm = sum(amps)
    n = 16
    pts = []
    for i in range(n):
        t = 2 * math.pi * i / n
        rr = 1 + wobble * sum(a * math.cos(k * t + ph) for (k, _), a, ph in zip(harm, amps, phases)) / norm
        pts.append((rr * math.cos(t), rr * math.sin(t)))
    return [Path([_smooth_closed(pts)])]


def _ticket_geo(w, h, notch=0.17, radius=0.07, style="sides", **_):
    """Ticket: rounded card with circular notches, on the two short sides or at the corners."""
    m = min(w, h)
    rc = radius * m
    if style == "sides":
        rn = notch * h
        pen = Pen().move(rc, 0).line(w - rc, 0).arc(w - rc, rc, rc, rc, -math.pi / 2, 0)
        pen.line(w, h / 2 - rn).arc(w, h / 2, rn, rn, -math.pi / 2, -1.5 * math.pi).line(w, h - rc)
        pen.arc(w - rc, h - rc, rc, rc, 0, math.pi / 2).line(rc, h)
        pen.arc(rc, h - rc, rc, rc, math.pi / 2, math.pi).line(0, h / 2 + rn)
        pen.arc(0, h / 2, rn, rn, math.pi / 2, -math.pi / 2).line(0, rc)
        pen.arc(rc, rc, rc, rc, math.pi, 1.5 * math.pi).close()
    elif style == "corners":
        rn = notch * m * 1.2
        pen = Pen().move(rn, 0).line(w - rn, 0).arc(w, 0, rn, rn, math.pi, math.pi / 2)
        pen.line(w, h - rn).arc(w, h, rn, rn, -math.pi / 2, -math.pi).line(rn, h)
        pen.arc(0, h, rn, rn, 0, -math.pi / 2).line(0, rn)
        pen.arc(0, 0, rn, rn, math.pi / 2, 0).close()
    else:
        raise ShapeError("ticket: style must be 'sides' or 'corners'")
    return [Path(pen.subs)]


def _bubble_bottom(w, h, pos, tail, radius, lean):
    """Speech bubble with its tail on the bottom side (the other sides are mirrors)."""
    td = tail * min(w, h)  # tail depth
    bh = h - td  # body height
    r = min(radius * min(w, bh), bh / 2, w / 2)
    bw = min(0.85 * td + 0.03 * w, 0.3 * w)  # tail base width
    lo, hi = r + 0.02 * w, w - r - bw - 0.02 * w
    b0 = lo + pos * max(0.0, hi - lo)
    b1 = b0 + bw
    tipx = b0 + bw / 2 + lean * (bw * 0.5 + 0.35 * td)
    pen = Pen().move(r, 0).line(w - r, 0).arc(w - r, r, r, r, -math.pi / 2, 0).line(w, bh - r)
    pen.arc(w - r, bh - r, r, r, 0, math.pi / 2).line(b1, bh)
    # both sides of the tail bow slightly inwards, which makes the point read as a point
    pen.curve(b1 + (tipx - b1) * 0.15, bh + td * 0.45, tipx + (b1 - tipx) * 0.18, h - td * 0.12, tipx, h)
    pen.curve(tipx + (b0 - tipx) * 0.30, h - td * 0.25, b0 + (tipx - b0) * 0.10, bh + td * 0.50, b0, bh)
    pen.line(r, bh).arc(r, bh - r, r, r, math.pi / 2, math.pi).line(0, r)
    pen.arc(r, r, r, r, math.pi, 1.5 * math.pi).close()
    return [Path(pen.subs)]


def _bubble_geo(w, h, side="bottom", pos=0.25, tail=None, radius=0.22, lean=-0.6, **_):
    """Speech bubble. ``side``: where the tail is; ``pos`` (0..1) along that side;
    ``tail``: tail depth over the shorter side of the box (default 0.22 on top or
    bottom, 0.3 on the left or right, where a short tail looks stunted);
    ``lean`` (-1..1): tip offset along the side."""
    if tail is None:
        tail = 0.3 if side in ("left", "right") else 0.22
    if not 0 <= pos <= 1 or not 0.1 <= tail <= 0.4 or not -1 <= lean <= 1:
        raise ShapeError("bubble: pos 0..1, tail 0.1..0.4, lean -1..1")
    if side == "bottom":
        return _bubble_bottom(w, h, pos, tail, radius, lean)
    if side == "top":
        return _map(_bubble_bottom(w, h, pos, tail, radius, lean), lambda x, y: (x, h - y))
    if side == "right":
        return _map(_bubble_bottom(h, w, pos, tail, radius, lean), lambda x, y: (y, x))
    if side == "left":
        return _map(_bubble_bottom(h, w, pos, tail, radius, lean), lambda x, y: (w - y, x))
    raise ShapeError("bubble: side must be top, bottom, left or right")


def _tear(rnd: random.Random, length: float, depth: float) -> List[Tuple[float, float]]:
    """Profile of a torn edge: (along, inward offset) points from 0 to length.
    A slow meander plus fine irregular teeth, the two scales of real torn paper."""
    pts = [(0.0, depth * 0.5)]
    x = 0.0
    phase, freq = rnd.random() * 2 * math.pi, 2 * math.pi / max(length * 0.45, depth * 6)
    while True:
        x += depth * (0.35 + 0.9 * rnd.random())
        if x >= length - depth * 0.3:
            break
        meander = 0.5 + 0.28 * math.sin(phase + x * freq)
        teeth = (rnd.random() - 0.5) * 0.55
        pts.append((x, depth * min(1.0, max(0.0, meander + teeth))))
    pts.append((length, depth * 0.5))
    return pts


def _torn_geo(w, h, edges="bottom", seed=0, depth=0.06, **_):
    """Paper card with one or more torn edges (``edges``: any of top, right, bottom, left,
    joined by '+', e.g. 'top+bottom'). ``depth``: tear depth over the box size across the edge."""
    wanted = set(edges.split("+"))
    if not wanted <= {"top", "right", "bottom", "left"}:
        raise ShapeError("torn: edges must be top, right, bottom or left, joined by '+'")
    rnd = random.Random(int(seed))
    pen = Pen().move(0, 0)

    def edge(name, length, across, to_point):
        if name not in wanted:
            pen.line(*to_point(length, 0.0))
            return
        d = min(depth * across, 0.25 * EMU_PER_INCH)
        for along, off in _tear(rnd, length, d)[1:]:
            pen.line(*to_point(along, off))

    edge("top", w, h, lambda a, o: (a, o))
    edge("right", h, w, lambda a, o: (w - o, a))
    edge("bottom", w, h, lambda a, o: (w - a, h - o))
    edge("left", h, w, lambda a, o: (o, h - a))
    sub = pen.close().subs[0]
    # a torn edge never ends exactly on the corner: drop duplicate points
    clean = [sub[0]] + [s for i, s in enumerate(sub[1:], 1) if s[0] != "L" or s[1:] != sub[i - 1][1:3]]
    return [Path([clean])]


def _polygon_geo(sides=6, rounding=0.12, **_):
    """Regular polygon, flat side down, corners slightly rounded (``rounding``: share of
    each side used by the round corner, 0 for sharp)."""
    if not 3 <= int(sides) <= 12:
        raise ShapeError("polygon: sides must be between 3 and 12")
    n = int(sides)
    start = math.pi / 2 + math.pi / n
    vs = [(math.cos(start + 2 * math.pi * i / n), math.sin(start + 2 * math.pi * i / n)) for i in range(n)]
    if rounding <= 0:
        pen = Pen().move(*vs[0])
        for v in vs[1:]:
            pen.line(*v)
        return [Path(pen.close().subs)]
    cut = min(rounding, 0.45)
    pen = Pen()
    for i in range(n):
        p, v, q = vs[i - 1], vs[i], vs[(i + 1) % n]
        a = (v[0] + (p[0] - v[0]) * cut, v[1] + (p[1] - v[1]) * cut)
        b = (v[0] + (q[0] - v[0]) * cut, v[1] + (q[1] - v[1]) * cut)
        if i == 0:
            pen.move(*a)
        else:
            pen.line(*a)
        k = 0.5523
        pen.curve(a[0] + (v[0] - a[0]) * k, a[1] + (v[1] - a[1]) * k, b[0] + (v[0] - b[0]) * k,
                  b[1] + (v[1] - b[1]) * k, *b)
    return [Path(pen.close().subs)]


def _sketch_side(pen: Pen, rnd: random.Random, p0, p1, amp: float, start: bool) -> None:
    """One side of a sketch from p0 to p1: a single slow bow (one or two bumps), the
    way a quick pen stroke wanders, never fine noise."""
    (x0, y0), (x1, y1) = p0, p1
    dx, dy = x1 - x0, y1 - y0
    ln = math.hypot(dx, dy) or 1.0
    nx, ny = -dy / ln, dx / ln
    b1, b2 = (rnd.random() - 0.5) * 2 * amp, (rnd.random() - 0.5) * 2 * amp
    t1, t2 = 0.25 + 0.15 * rnd.random(), 0.6 + 0.15 * rnd.random()
    if start:
        pen.move(x0, y0)
    pen.curve(x0 + dx * t1 + nx * b1, y0 + dy * t1 + ny * b1, x0 + dx * t2 + nx * b2, y0 + dy * t2 + ny * b2, x1, y1)


def _handdrawn_geo(w, h, seed=0, jitter=0.03, strokes=2, overshoot=0.5, **_):
    """Hand-drawn card, sketched with a pen. The filled outline has four slightly crooked
    corners joined by slow bows. With ``strokes=2`` a second pencil pass, outline only,
    draws each side again as a separate stroke: offset from the first, at a slightly
    different angle (not parallel) and running past the corners (``overshoot``, 0..1),
    as quick sketches do. Where the two passes meet the line looks thicker.
    ``jitter``: wobble amplitude over the shorter side (0..0.05)."""
    if strokes not in (1, 2):
        raise ShapeError("handdrawn: strokes must be 1 or 2")
    if not 0 <= jitter <= 0.05:
        raise ShapeError("handdrawn: jitter must be between 0 and 0.05")
    if not 0 <= overshoot <= 1:
        raise ShapeError("handdrawn: overshoot must be between 0 and 1")
    rnd = random.Random(int(seed))
    m = min(w, h)
    amp = jitter * m + 0.004 * max(w, h)
    over = overshoot * 3 * amp  # how far the second pass runs past a corner
    mg = max(over + amp, 2.8 * amp)  # margin that keeps every stroke inside the box

    def corner(x, y):
        return x + (rnd.random() - 0.5) * 1.6 * amp, y + (rnd.random() - 0.5) * 1.6 * amp

    cs = [corner(mg, mg), corner(w - mg, mg), corner(w - mg, h - mg), corner(mg, h - mg)]
    pen = Pen()
    for i in range(4):
        _sketch_side(pen, rnd, cs[i], cs[(i + 1) % 4], amp, i == 0)
    geom = [Path(pen.close().subs)]
    if strokes == 2:
        pen2 = Pen()
        for i in range(4):
            (ax, ay), (bx, by) = cs[i], cs[(i + 1) % 4]
            dx, dy = bx - ax, by - ay
            ln = math.hypot(dx, dy) or 1.0
            ux, uy, nx, ny = dx / ln, dy / ln, -dy / ln, dx / ln
            # each end gets its own offset, so the two passes are never parallel
            o0 = (0.4 + 0.8 * rnd.random()) * amp * (1 if rnd.random() < 0.5 else -1)
            o1 = (0.4 + 0.8 * rnd.random()) * amp * (1 if rnd.random() < 0.5 else -1)
            e0 = over * (0.5 + 0.5 * rnd.random())
            e1 = over * (0.5 + 0.5 * rnd.random())
            p0 = (ax - ux * e0 + nx * o0, ay - uy * e0 + ny * o0)
            p1 = (bx + ux * e1 + nx * o1, by + uy * e1 + ny * o1)
            _sketch_side(pen2, rnd, p0, p1, amp * 0.8, True)
        geom.append(Path(pen2.subs, fill=False, stroke=True))
    return geom


# ---------------------------------------------------------------- icons
def icon_geometry(svg_path) -> Geometry:
    """Filled SVG icon (one colour) as design-space geometry, cropped to its ink."""
    import sys

    here = str(_FsPath(__file__).resolve().parent)
    if here not in sys.path:
        sys.path.insert(0, here)
    from _svg import parse_svg  # vendored converter, same author (see its header)

    text = _FsPath(svg_path).read_text(encoding="utf-8")
    _, items, _ = parse_svg(text)
    subs: List[Sub] = []
    for it in items:
        if it["fill"] == "none":
            continue
        cur: Sub = []
        for s in it["segs"]:
            if s[0] == "M":
                if cur:
                    subs.append(cur)
                cur = [s]
            elif not cur:
                continue
            else:
                cur.append(s)
        if cur:
            subs.append(cur)
    if not subs:
        raise ShapeError(f"icon: nothing filled in {svg_path}")
    return [Path(subs)]


def _icon_geo(svg=None, **_):
    if svg is None:
        raise ShapeError("icon: pass svg=<path to a filled SVG icon>")
    return icon_geometry(svg)


# --------------------------------------------------------------- registry
class Kind(NamedTuple):
    build: Callable[..., Geometry]
    box: bool  # True: drawn at the box size; False: design space, stretched to the box
    aspect: Optional[Tuple[float, float]]  # fixed (min, max) of w / h; None = natural x tolerance
    tolerance: float = 1.25  # allowed stretch from the natural proportion (design shapes)
    pad: float = 0.04  # breathing room around the text box, share of the shorter side
    text_min: float = 0.0  # narrowest text box allowed (w / h): keeps text boxes usable in tall shapes
    min_size: Optional[Tuple[float, float]] = None  # smallest usable (w, h) in inches for 14 pt text
    max_words: Optional[int] = None  # most words the shape should ever hold, whatever its size


REFERENCE_PT = 14.0  # font size that min_size is given for; it scales with the font size

KINDS: Dict[str, Kind] = {
    "rectangle": Kind(_rect_geo, True, (0.1, 10.0), pad=0.06),
    "rounded": Kind(_rounded_geo, True, (0.1, 10.0), pad=0.06),
    "pill": Kind(_pill_geo, True, (1.8, 12.0)),
    "ellipse": Kind(_ellipse_geo, True, (0.5, 2.5)),
    "circle": Kind(_circle_geo, True, (0.8, 1.25)),
    # a drop holds a few words, never a sentence: below 2 in (at 14 pt) words break
    "drop": Kind(_drop_geo, False, None, 1.2, text_min=1.2, min_size=(1.2, 2.0), max_words=3),
    "cloud": Kind(_cloud_geo, False, None, 1.3),
    "blob": Kind(_blob_geo, False, None, 1.6, text_min=0.9),
    "ticket": Kind(_ticket_geo, True, (1.3, 5.0)),
    "bubble": Kind(_bubble_geo, True, (0.6, 4.0)),
    "torn": Kind(_torn_geo, True, (0.25, 6.0)),
    "polygon": Kind(_polygon_geo, False, None, 1.2, text_min=0.9),
    "handdrawn": Kind(_handdrawn_geo, True, (0.2, 8.0), pad=0.06),
    "icon": Kind(_icon_geo, False, None, 1.25, text_min=0.8),
}


def kinds() -> List[str]:
    return list(KINDS)


def _kind(kind: str) -> Kind:
    try:
        return KINDS[kind]
    except KeyError:
        raise ShapeError(f"unknown shape '{kind}' (known: {', '.join(KINDS)})") from None


def natural_aspect(kind: str, **params) -> Optional[float]:
    """w / h at which a design-space shape is drawn undistorted (None for box shapes)."""
    k = _kind(kind)
    if k.box:
        return None
    x0, y0, x1, y1 = bbox(k.build(**params))
    return (x1 - x0) / (y1 - y0)


def aspect_bounds(kind: str, **params) -> Tuple[float, float]:
    """(min, max) of w / h that the shape accepts without visible distortion."""
    k = _kind(kind)
    if k.aspect is not None:
        if kind == "bubble" and params.get("side") in ("left", "right"):
            return 0.9, 5.0
        return k.aspect
    nat = natural_aspect(kind, **params)
    return nat / k.tolerance, nat * k.tolerance


def _aspect_ok(kind: str, w: float, h: float, **params) -> bool:
    if w <= 0 or h <= 0:
        return False
    lo, hi = aspect_bounds(kind, **params)
    return lo - 1e-9 <= w / h <= hi + 1e-9


def min_size(kind: str, size_pt: float = REFERENCE_PT) -> Optional[Tuple[int, int]]:
    """Smallest usable (w, h) in EMU of a shape for text at ``size_pt``, or None when
    the kind has no minimum. Below it words would break mid-word."""
    ms = _kind(kind).min_size
    if ms is None:
        return None
    f = size_pt / REFERENCE_PT
    return int(round(ms[0] * EMU_PER_INCH * f)), int(round(ms[1] * EMU_PER_INCH * f))


def fits(kind: str, w: float, h: float, *, size_pt: float = REFERENCE_PT, text: Optional[str] = None,
         char_em: float = 0.52, caps: bool = False, **params) -> bool:
    """True when the shape can be used at a w x h box (EMU) for text at ``size_pt``:
    the box keeps the shape within its aspect bounds and is at least the kind's
    ``min_size``. With ``text``, the whole text must also fit the safe text box with
    whole words only: no word wider than a line, no more words than ``word_limit``."""
    if not _aspect_ok(kind, w, h, **params):
        return False
    ms = min_size(kind, size_pt)
    if ms is not None and (w < ms[0] - 1 or h < ms[1] - 1):
        return False
    if text is None:
        return True
    words = text.split()
    if not words:
        return True
    try:
        n = word_limit(kind, (w, h), size_pt, text=text, char_em=char_em, caps=caps, **params)
    except ShapeError:
        return False
    return n >= len(words)


def outline(kind: str, w: float, h: float, *, strict: bool = True, **params) -> Geometry:
    """The geometry of a shape sized w x h (EMU), as Path tuples."""
    k = _kind(kind)
    if strict and not _aspect_ok(kind, w, h, **params):
        lo, hi = aspect_bounds(kind, **params)
        raise ShapeError(f"{kind}: box {w / h:.2f}:1 outside its aspect bounds {lo:.2f} to {hi:.2f}")
    if k.box:
        return k.build(w, h, **params)
    return _fit(k.build(**params), w, h)


# ------------------------------------------------------------ safe box
RASTER = 200  # samples along the longer side of the box


def _mask(geom: Geometry, w: float, h: float):
    from PIL import Image, ImageChops, ImageDraw, ImageFilter

    s = RASTER / max(w, h)
    if min(w, h) * s < 40:  # very long boxes: keep at least 40 samples across
        s = 40 / min(w, h)
    pw, ph = max(1, int(math.ceil(w * s))), max(1, int(math.ceil(h * s)))
    mask = Image.new("1", (pw, ph), 0)
    for p in geom:
        if not p.fill:
            continue
        for sub in p.subs:
            poly = [(x * s, y * s) for x, y in flatten(sub)]
            if len(poly) < 3:
                continue
            layer = Image.new("1", (pw, ph), 0)
            ImageDraw.Draw(layer).polygon(poly, fill=1)
            mask = ImageChops.logical_xor(mask, layer)  # even-odd: holes stay holes
    # a sample counts as inside only when its neighbours are too: the rectangle then
    # stays inside the outline, not just inside the sampled pixels
    mask = mask.convert("L").filter(ImageFilter.MinFilter(3))
    return mask, s


def _largest_rect(mask, min_ratio: float = 0.0) -> Tuple[int, int, int, int]:
    """Largest all-inside rectangle (x0, y0, x1, y1) in pixel cells, end exclusive, at
    least ``min_ratio`` wide for its height. Ties go to the rectangle closest to the
    centre, so the result is deterministic.

    Every maximal rectangle of the histogram scan is a candidate; when it is too narrow
    it is shortened (any part of an inside rectangle is inside), centred on its middle."""
    pw, ph = mask.size
    data = mask.tobytes()
    heights = [0] * pw
    best, best_key = (0, 0, 0, 0), (0, 0.0)
    cxm, cym = pw / 2, ph / 2
    for y in range(ph):
        row = data[y * pw:(y + 1) * pw]
        for x in range(pw):
            heights[x] = heights[x] + 1 if row[x] else 0
        stack: List[int] = []
        for x in range(pw + 1):
            hx = heights[x] if x < pw else 0
            while stack and heights[stack[-1]] >= hx:
                top = stack.pop()
                full = heights[top]
                left = stack[-1] + 1 if stack else 0
                width = x - left
                height = full
                if min_ratio > 0:
                    height = min(full, int(width / min_ratio))
                area = height * width
                if area and area >= best_key[0]:
                    x0 = left
                    y0 = y - full + 1 + (full - height) // 2
                    dist = -abs((x0 + x) / 2 - cxm) - abs(y0 + height / 2 - cym)
                    key = (area, dist)
                    if key > best_key:
                        best_key, best = key, (x0, y0, x, y0 + height)
            stack.append(x)
    return best


def safe_insets(geom: Geometry, w: float, h: float, pad: float = 0.04, text_min: float = 0.0) -> Insets:
    """Largest inscribed axis-aligned rectangle of the filled geometry (at least
    ``text_min`` wide for its height), minus ``pad`` (share of the shorter side) all
    round, as insets from the box."""
    mask, s = _mask(geom, w, h)
    x0, y0, x1, y1 = _largest_rect(mask, text_min)
    if x1 <= x0 or y1 <= y0:
        raise ShapeError("no room for text inside this shape")
    p = pad * min(w, h)
    left, top = int(math.ceil(x0 / s + p)), int(math.ceil(y0 / s + p))
    right, bottom = int(math.ceil(w - x1 / s + p)), int(math.ceil(h - y1 / s + p))
    right, bottom = max(0, right), max(0, bottom)
    if left + right >= w or top + bottom >= h:
        raise ShapeError("no room for text inside this shape")
    return Insets(left, top, right, bottom)


def char_limit(w_emu: float, h_emu: float, size_pt: float, char_em: float = 0.52, caps: bool = False) -> int:
    """Characters that fit a text box, with the style contract formula:
    per line = floor(width pt / (size x char_em)), lines = floor(height pt / (size x 1.2)),
    limit = max(4, floor(per line x lines x 0.85)); char_em x 1.2 for text in capitals.
    A budget only: cut text with ``cut_words`` or ``fit_words``, never mid-word."""
    em = char_em * (1.2 if caps else 1.0)
    per_line = math.floor((w_emu / EMU_PER_PT) / (size_pt * em))
    lines = math.floor((h_emu / EMU_PER_PT) / (size_pt * 1.2))
    return max(4, math.floor(per_line * lines * 0.85))


def text_limit(kind: str, box, size_pt: float, char_em: float = 0.52, caps: bool = False, **params) -> int:
    """Character limit of the safe text box of a shape at a box (x, y, w, h) or (w, h)."""
    w, h = _wh(box)
    ins = make_shape(kind, box, **params).insets
    return char_limit(w - ins.left - ins.right, h - ins.top - ins.bottom, size_pt, char_em, caps)


ELLIPSIS = "\u2026"
WORD_SAFETY = 1.15  # a single word is measured wider than average text: no average to lean on


def cut_words(text: str, limit: int) -> str:
    """``text`` cut to at most ``limit`` characters at a word boundary, with an
    ellipsis when something was cut. Never breaks a word: returns "" when even the
    first word (plus the ellipsis) is too long."""
    words = text.split()
    if len(" ".join(words)) <= limit:
        return " ".join(words)
    out = ""
    for wd in words:
        cand = (out + " " + wd) if out else wd
        if len(cand.rstrip(",;:.")) + 1 > limit:
            break
        out = cand
    return out.rstrip(",;:.") + ELLIPSIS if out else ""


def _wrap_count(words: Sequence[str], w_emu: float, h_emu: float, size_pt: float, char_em: float,
                caps: bool) -> int:
    """How many leading words a greedy line wrap places in the box without breaking one."""
    em = char_em * (1.2 if caps else 1.0)
    width = w_emu / EMU_PER_PT
    lines = math.floor((h_emu / EMU_PER_PT) / (size_pt * 1.2))
    unit = size_pt * em

    def wide(s):
        return len(s) * unit

    line, used, n = "", 1, 0
    if lines < 1:
        return 0
    for wd in words:
        if wide(wd.rstrip(ELLIPSIS)) * WORD_SAFETY + wide(wd) - wide(wd.rstrip(ELLIPSIS)) > width:
            break  # this word would have to be broken (a trailing ellipsis is narrow, no margin)
        cand = (line + " " + wd) if line else wd
        if wide(cand) <= width:
            line = cand
        else:
            if used == lines:
                break
            used, line = used + 1, wd
        n += 1
    return n


def _safe_wh(kind: str, box, **params) -> Tuple[int, int]:
    w, h = _wh(box)
    ins = make_shape(kind, box, strict=False, **params).insets
    return w - ins.left - ins.right, h - ins.top - ins.bottom


def word_limit(kind: str, box, size_pt: float, *, text: Optional[str] = None, char_em: float = 0.52,
               caps: bool = False, **params) -> int:
    """Whole words the safe text box of a shape holds at ``size_pt``, capped by the
    kind's ``max_words`` and by the character limit.

    With ``text``: how many of its leading words fit (greedy line wrap, no word ever
    broken; 0 when the first word is wider than a line). Without: an estimate for
    average words of 6 characters (5 letters and a space). The restyler refuses text
    with more words than this rather than letting PowerPoint break one."""
    tw, th = _safe_wh(kind, box, **params)
    cap = _kind(kind).max_words
    limit = char_limit(tw, th, size_pt, char_em, caps)
    if text is None:
        words = ["abcde"] * 200
    else:
        words = text.split()
    n = _wrap_count(words, tw, th, size_pt, char_em, caps)
    while n and len(" ".join(words[:n])) > limit:
        n -= 1
    return min(n, cap) if cap is not None else n


def fit_words(kind: str, box, text: str, size_pt: float, *, char_em: float = 0.52, caps: bool = False,
              **params) -> str:
    """``text`` cut to what the safe text box of the shape holds: whole words only, an
    ellipsis when cut, "" when not even one word fits."""
    words = text.split()
    n = word_limit(kind, box, size_pt, text=text, char_em=char_em, caps=caps, **params)
    if n >= len(words):
        return " ".join(words)
    tw, th = _safe_wh(kind, box, **params)
    while n > 0:  # the ellipsis must fit too
        out = " ".join(words[:n]).rstrip(",;:.") + ELLIPSIS
        if _wrap_count(out.split(), tw, th, size_pt, char_em, caps) == len(out.split()):
            return out
        n -= 1
    return ""


# --------------------------------------------------------------- output
def _q(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


def _wh(box) -> Tuple[int, int]:
    if len(box) == 4:
        w, h = box[2], box[3]
    elif len(box) == 2:
        w, h = box
    else:
        raise ShapeError("box must be (x, y, w, h) or (w, h)")
    w, h = int(round(w)), int(round(h))
    if w <= 0 or h <= 0:
        raise ShapeError("box width and height must be positive")
    return w, h


def custgeom(geom: Geometry, w: int, h: int, insets: Optional[Insets] = None) -> etree._Element:
    """Build the a:custGeom element. The path space is the box in EMU; the text
    rectangle follows the shape when it is resized (guides proportional to w and h)."""
    root = etree.Element(_q("custGeom"), nsmap={"a": A_NS})
    etree.SubElement(root, _q("avLst"))
    gd = etree.SubElement(root, _q("gdLst"))
    etree.SubElement(root, _q("ahLst"))
    etree.SubElement(root, _q("cxnLst"))
    rect = etree.SubElement(root, _q("rect"))
    if insets is None:
        rect.attrib.update({"l": "l", "t": "t", "r": "r", "b": "b"})
    else:
        vals = {"txL": (insets.left, w, "w"), "txT": (insets.top, h, "h"),
                "txR": (w - insets.right, w, "w"), "txB": (h - insets.bottom, h, "h")}
        for name, (v, size, ref) in vals.items():
            etree.SubElement(gd, _q("gd"), name=name, fmla=f"*/ {ref} {int(round(v * 100000 / size))} 100000")
        rect.attrib.update({"l": "txL", "t": "txT", "r": "txR", "b": "txB"})
    plst = etree.SubElement(root, _q("pathLst"))

    def pt(parent, x, y):
        etree.SubElement(parent, _q("pt"), x=str(int(round(x))), y=str(int(round(y))))

    for p in geom:
        attrs = {"w": str(w), "h": str(h)}
        if not p.fill:
            attrs["fill"] = "none"
        if not p.stroke:
            attrs["stroke"] = "0"
        el = etree.SubElement(plst, _q("path"), **attrs)
        for sub in p.subs:
            for s in sub:
                if s[0] == "M":
                    pt(etree.SubElement(el, _q("moveTo")), s[1], s[2])
                elif s[0] == "L":
                    pt(etree.SubElement(el, _q("lnTo")), s[1], s[2])
                elif s[0] == "C":
                    c = etree.SubElement(el, _q("cubicBezTo"))
                    for k in (1, 3, 5):
                        pt(c, s[k], s[k + 1])
                else:
                    etree.SubElement(el, _q("close"))
    return root


def make_shape(kind: str, box, *, strict: bool = True, **params) -> ShapeResult:
    """Build shape ``kind`` for ``box`` = (x, y, w, h) or (w, h), in EMU.

    Returns (geom, insets): the a:custGeom element, and the safe text box as insets
    from the box. ``strict`` refuses boxes outside the shape's aspect bounds.
    Parameters per kind: see the builders (``radius``, ``tilt``, ``seed``, ``side``...);
    any kind also takes ``pad`` (room around the text box, share of the shorter side)
    and ``text_min`` (narrowest text box, w / h).
    """
    w, h = _wh(box)
    geom = outline(kind, w, h, strict=strict, **params)
    k = _kind(kind)
    insets = safe_insets(geom, w, h, params.get("pad", k.pad), params.get("text_min", k.text_min))
    return ShapeResult(custgeom(geom, w, h, insets), insets)


def icon_as_container(svg_path, box=(EMU_PER_INCH, EMU_PER_INCH), **params) -> ShapeResult:
    """A filled SVG icon (Phosphor fill, Tabler filled...) as a container silhouette."""
    return make_shape("icon", box, svg=str(svg_path), **params)


def set_geometry(shape, geom: etree._Element) -> None:
    """Replace the geometry of a python-pptx autoshape with an a:custGeom element."""
    sppr = shape._element.spPr
    old = sppr.find(_q("prstGeom"))
    if old is None:
        old = sppr.find(_q("custGeom"))
    if old is not None:
        old.addprevious(geom)
        sppr.remove(old)
    else:
        xfrm = sppr.find(_q("xfrm"))
        (xfrm.addnext if xfrm is not None else lambda e: sppr.insert(0, e))(geom)


def add_shape(slide, kind: str, box, *, name: Optional[str] = None, strict: bool = True, **params):
    """Add shape ``kind`` to a python-pptx slide at box (x, y, w, h). Returns (shape, insets).
    The shape's own text frame already uses the safe text box. It is flat (no theme
    style, so no shadow), filled with theme accent1 and has no outline."""
    from pptx.enum.shapes import MSO_SHAPE

    x, y, w, h = (int(round(v)) for v in box)
    geom, insets = make_shape(kind, (x, y, w, h), strict=strict, **params)
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    set_geometry(shp, geom)
    # flat and explicit: drop the default shape style (its effect reference draws a
    # shadow in LibreOffice even under an empty effect list); theme fill accent1, no
    # outline, theme text colour, until the caller sets its own
    from pptx.enum.dml import MSO_THEME_COLOR

    style = shp._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        shp._element.remove(style)
    shp.fill.solid()
    shp.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    shp.line.fill.background()
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    if name:
        shp.name = name
    return shp, insets
