"""Ambience layer: scope, z-order, theme colours, determinism, contrast gate, deck mode."""
import json
import shutil
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

import _ambience as AMB
import restyle as R

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "make-clean-1.1.0-ambience"
IN_SCOPE = {"L-02", "L-09", "L-12", "L-14"}  # covers, dividers, statements of density Low
OUT_OF_SCOPE = {"L-24", "L-37", "L-41"}  # medium text, a KPI row, a table
A = AMB.A


def style_with(ambience, tmp_path, name="amb"):
    p = tmp_path / f"{name}.json"
    p.write_text(json.dumps({"name": name, "seed": 5, "roles": {}, "ambience": ambience}))
    return R.load_style(str(p))


def slides_by_id(pack: Path, lib="layouts"):
    _, rows = R.read_index(pack / f"{lib}.csv")
    prs = Presentation(str(pack / f"{lib}.pptx"))
    ids = {int(r["slide"]): r["id"] for r in rows}
    return {ids[n]: s for n, s in enumerate(prs.slides, 1)}


def ambience_group(slide):
    return [sh for sh in slide.shapes if sh.name.startswith(AMB.GROUP_NAME)]


@pytest.fixture(scope="module", params=AMB.MOTIFS)
def motif_pack(request, tmp_path_factory):
    tmp = tmp_path_factory.mktemp(request.param)
    style = style_with({"motif": request.param}, tmp, request.param)
    out, report = R.run(FIXTURE, style, tmp / "pack")
    return request.param, out, report


def test_motif_on_low_text_slides_only(motif_pack):
    motif, out, report = motif_pack
    items = {it["id"]: it for it in report["libraries"]["layouts"]["ambience"]["items"]}
    assert {i for i, it in items.items() if it["status"] == "applied"} == IN_SCOPE
    assert all(items[i]["status"] == "out-of-scope" for i in OUT_OF_SCOPE)
    slides = slides_by_id(out)
    for row_id, slide in slides.items():
        groups = ambience_group(slide)
        if row_id in IN_SCOPE:
            assert len(groups) == 1 and groups[0].name == f"{AMB.GROUP_NAME} > {motif}"
        else:
            assert not groups
    visuals = report["libraries"]["visuals"]["ambience"]["items"]
    assert visuals and all(it["status"] == "out-of-scope" for it in visuals)
    for slide in slides_by_id(out, "visuals").values():
        assert not ambience_group(slide)


def test_motif_behind_everything(motif_pack):
    _, out, _ = motif_pack
    for row_id, slide in slides_by_id(out).items():
        if row_id in IN_SCOPE:
            # first shape of the tree = drawn first = behind every text and rule
            assert slide.shapes[0].name.startswith(AMB.GROUP_NAME)
            assert all(not sh.name.startswith(AMB.GROUP_NAME) for sh in list(slide.shapes)[1:])


def test_theme_colours_decorative_and_editable(motif_pack):
    _, out, _ = motif_pack
    for row_id, slide in slides_by_id(out).items():
        if row_id not in IN_SCOPE:
            continue
        grp = ambience_group(slide)[0]._element
        assert not list(grp.iter(A + "srgbClr")), "motif colours must be theme references"
        clrs = list(grp.iter(A + "schemeClr"))
        assert clrs and all(c.find(A + "alpha") is not None for c in clrs)
        assert int(clrs[0].find(A + "alpha").get("val")) < 50000
        assert grp.find(f".//{{{AMB.ADEC}}}decorative") is not None
        assert list(grp.iter(A + "custGeom")) and not list(grp.iter(A + "blip")), "native shapes, no bitmap"
        # dark slides get the light slot, light slides the dark one
        dark = row_id in ("L-02", "L-09")
        assert clrs[0].get("val") == ("bg1" if dark else "tx2")


def test_excluded_slides_unchanged(motif_pack, tmp_path):
    _, out, _ = motif_pack
    plain, _ = R.run(FIXTURE, style_with(None, tmp_path, "plain"), tmp_path / "plain")
    a, b = slides_by_id(plain), slides_by_id(out)
    for row_id in OUT_OF_SCOPE:
        assert etree.tostring(a[row_id]._element) == etree.tostring(b[row_id]._element)


def test_deterministic_and_seeded(tmp_path):
    style = style_with({"motif": "rain"}, tmp_path)
    one, _ = R.run(FIXTURE, style, tmp_path / "one")
    two, _ = R.run(FIXTURE, style, tmp_path / "two")
    assert (one / "layouts.pptx").read_bytes() == (two / "layouts.pptx").read_bytes()
    other, _ = R.run(FIXTURE, style_with({"motif": "rain", "seed": 3}, tmp_path, "other"), tmp_path / "three")
    a = etree.tostring(ambience_group(slides_by_id(one)["L-12"])[0]._element)
    b = etree.tostring(ambience_group(slides_by_id(other)["L-12"])[0]._element)
    assert a != b


def test_rain_and_mist_keep_clear_of_text(tmp_path):
    out, _ = R.run(FIXTURE, style_with({"motif": "rain"}, tmp_path), tmp_path / "rain")
    for row_id, slide in slides_by_id(out).items():
        if row_id not in IN_SCOPE:
            continue
        texts = [AMB._box(sh) for sh in slide.shapes if AMB._text(sh).strip()]
        for sp in ambience_group(slide)[0].shapes:
            x0, y0 = AMB._box(sp)[:2]
            pts = [(x0 + int(p.get("x")), y0 + int(p.get("y"))) for p in sp._element.iter(A + "pt")]
            for (ax, ay), (bx, by) in zip(pts[::2], pts[1::2]):
                seg = AMB._seg_box(ax, ay, bx, by)
                assert not any(AMB._overlap(seg, t) for t in texts)


def test_contrast_gate_unit():
    white, grey, black = (1.0, 1.0, 1.0), (0.4, 0.4, 0.4), (0.0, 0.0, 0.0)
    run = {"box": (0, 0, 100, 100), "rgb": grey, "need": 4.5}
    assert AMB.contrast(grey, white) > 4.5
    faint = [{"kind": "blob", "box": (50, 50, 100, 100), "alpha": 0.03}]
    strong = [{"kind": "blob", "box": (50, 50, 100, 100), "alpha": 0.5}]
    away = [{"kind": "blob", "box": (500, 500, 100, 100), "alpha": 0.9}]
    assert AMB.gate(faint, [run], black, white) == (faint, 0)
    assert AMB.gate(strong, [run], black, white) == ([], 1)
    assert AMB.gate(away, [run], black, white) == (away, 0)
    # overlapping blobs add up: after a few faint layers the next one tips the run below 4.5:1
    layers = [{"kind": "blob", "box": (0, 0, 100, 100), "alpha": 0.06} for _ in range(8)]
    kept, dropped = AMB.gate(layers, [run], black, white)
    assert dropped and kept
    a = 1.0
    for p in kept:
        a *= 1 - p["alpha"]
    assert AMB.contrast(grey, AMB.blend(black, 1 - a, white)) >= 4.5
    # lines side by side do not add up: each crosses the run on its own
    lines = [{"kind": "seg", "box": (10 * i, 0, 1, 100), "alpha": 0.06} for i in range(8)]
    assert AMB.gate(lines, [run], black, white) == (lines, 0)
    # a run already below its threshold may not get noticeably worse
    pale = {"box": (0, 0, 100, 100), "rgb": (0.7, 0.7, 0.7), "need": 4.5}
    assert AMB.gate(faint, [pale], black, white)[1] == 1


def test_contrast_gate_on_slides(tmp_path, monkeypatch):
    dens = json.loads(json.dumps(AMB.DENSITY))
    dens["low"]["grid"]["alpha"] = 0.9  # far too strong: the gate must cut the cells under text
    monkeypatch.setattr(AMB, "DENSITY", dens)
    _, report = R.run(FIXTURE, style_with({"motif": "grid"}, tmp_path), tmp_path / "loud")
    items = [it for it in report["libraries"]["layouts"]["ambience"]["items"] if it["id"] in IN_SCOPE]
    assert any(it["dropped_by_contrast"] for it in items)
    out = tmp_path / "loud"
    for row_id, slide in slides_by_id(out).items():
        if row_id not in IN_SCOPE:
            continue
        theme = AMB.Theme(slide)
        bg = theme.background()
        runs = AMB.text_runs(slide, theme)
        for sp in ambience_group(slide)[0].shapes:
            clr = next(sp._element.iter(A + "schemeClr"))
            alpha = int(clr.find(A + "alpha").get("val")) / 100000
            under = AMB.blend(theme.scheme(clr.get("val")), alpha, bg)
            x0, y0 = AMB._box(sp)[:2]
            pts = [(x0 + int(p.get("x")), y0 + int(p.get("y"))) for p in sp._element.iter(A + "pt")]
            for (ax, ay), (bx, by) in zip(pts[::2], pts[1::2]):
                seg = AMB._seg_box(ax, ay, bx, by)
                for r in runs:
                    if AMB._overlap(seg, r["box"]):
                        base = AMB.contrast(r["rgb"], bg)
                        need = r["need"] if base >= r["need"] else base * (1 - AMB.WORSE)
                        assert AMB.contrast(r["rgb"], under) >= need - 1e-9, f"{row_id}: a line under hard-to-read text"


def test_defaults_add_nothing(tmp_path):
    style = style_with(None, tmp_path)
    assert style["ambience"] is None
    out, report = R.run(FIXTURE, style, tmp_path / "pack")
    assert "ambience" not in report["libraries"]["layouts"]
    assert not any(ambience_group(s) for s in slides_by_id(out).values())


def test_bad_ambience_refused(tmp_path):
    for bad in ({"motif": "confetti"}, {"motif": "mist", "density": "loud"}, {"motif": "mist", "colour": "red"},
                {"motif": "mist", "sparkle": 1}, {"typography": {"major": ""}}, 42):
        with pytest.raises(R.RestyleError):
            style_with(bad, tmp_path)
    assert AMB.normalise("mist")["motif"] == "mist"


def test_reference_styles_carry_one_motif():
    assert R.load_style("organic")["ambience"]["motif"] == "mist"
    assert R.load_style("paper")["ambience"]["motif"] == "grid"


def test_deck_mode(tmp_path):
    deck = tmp_path / "deck.pptx"
    shutil.copyfile(FIXTURE / "layouts.pptx", deck)
    out, report = R.run(deck, style_with({"motif": "mist"}, tmp_path))
    ids = {int(r["slide"]): r["id"] for r in R.read_index(FIXTURE / "layouts.csv")[1]}
    items = {ids[int(it["id"].split()[1])]: it for it in report["ambience"]["items"]}
    assert {i for i, it in items.items() if it["status"] == "applied"} == IN_SCOPE
    assert items["L-41"]["reason"] == "slide with a chart or a table"
    prs = Presentation(str(out))
    for n, slide in enumerate(prs.slides, 1):
        assert bool(ambience_group(slide)) == (ids[n] in IN_SCOPE)


def test_typography_pair(tmp_path):
    style = style_with({"typography": {"major": "Georgia", "minor": "Verdana"}}, tmp_path)
    out, report = R.run(FIXTURE, style, tmp_path / "pack")
    assert report["libraries"]["layouts"]["ambience"]["typography"]["major"] == "Georgia"
    prs = Presentation(str(out / "layouts.pptx"))
    theme = etree.fromstring(prs.slide_masters[0].part.part_related_by(AMB.RT_THEME).blob)
    assert theme.find(f".//{A}majorFont/{A}latin").get("typeface") == "Georgia"
    assert theme.find(f".//{A}minorFont/{A}latin").get("typeface") == "Verdana"
    assert not any(ambience_group(s) for s in slides_by_id(out).values())  # fonts only, no motif
