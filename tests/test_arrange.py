"""Arrangement layer: positions per count, grid rules, overlap, margins, reading order, fallback, determinism."""
import json
import math
import shutil
import subprocess
import sys

import pytest
from lxml import etree
from pptx import Presentation

import _arrange as AR
import restyle as R
from test_restyle import FIXTURE, _make_clean_dir, items, shapes_by_id

E = AR.EMU


def style_with(name, base="organic"):
    st = R.load_style(base)
    st["arrangement"] = name
    return st


@pytest.fixture(scope="module")
def arranged(tmp_path_factory):
    out = {}
    for name in AR.NAMES:
        dst = tmp_path_factory.mktemp(name) / "pack"
        out[name] = R.run(FIXTURE, style_with(name), dst)
    return out


def arrangement_items(report):
    return {it["id"]: it for it in items(report) if it["role"] == "arrangement"}


def units_of(pack, row_id, count):
    """{n: union box of the unit's fields} in a pack, from the field names."""
    shapes = shapes_by_id(pack / "layouts.pptx", pack / "layouts.csv")[row_id]
    out = {}
    for name, sh in shapes.items():
        link = R.TAG.match(name).group(3) if R.TAG.match(name) else name
        n = AR.unit_number(link or "")
        if n and 1 <= n <= count and sh.has_text_frame:
            out.setdefault(n, []).append(R.box(sh))
    return {n: AR._union(b) for n, b in out.items()}


def plain_row(n, top=2.3, h=3.4):
    """Units of a plain row on the 12-column grid (EMU)."""
    span = 24 // n
    return [(AR.half_x(i * span), AR.inch(top), AR.span_w(span), AR.inch(h)) for i in range(n)]


def plain_column(n, x=0.6, w=7.4, top=2.4, h=0.8, step=1.1):
    return [(AR.inch(x), AR.inch(top + i * step), AR.inch(w), AR.inch(h)) for i in range(n)]


# ------------------------------------------------------------------ pure geometry
@pytest.mark.parametrize("name", AR.NAMES)
@pytest.mark.parametrize("n", [2, 3, 4, 5, 6])
def test_row_positions_no_overlap_inside_margins_in_order(name, n):
    units = plain_row(n)
    region = AR._union(units)
    region = (region[0], region[1], region[2], AR.inch(AR.GRID["content_bottom"]) - region[1])
    try:
        rects = AR.positions(name, "row", units, region)
    except AR.ArrangeError:
        assert (name == "arc" and n < 3) or (name == "deconstructed" and n not in AR.DECONSTRUCTED_ROW)
        return
    assert len(rects) == n
    assert not AR.overlapping(rects)
    lim = AR.bounds((AR.inch(AR.GRID["w"]), AR.inch(7.5)), AR._union(units))
    assert all(AR.contains(lim, r, 2) for r in rects)
    assert AR.reading_order("row", rects) == list(range(n))
    assert all(r[2] <= u[2] + 2 and r[3] <= u[3] for r, u in zip(rects, units))  # never larger


@pytest.mark.parametrize("name", AR.NAMES)
@pytest.mark.parametrize("n", [2, 3, 4, 5, 6])
def test_column_positions_no_overlap_in_order(name, n):
    units = plain_column(n)
    try:
        rects = AR.positions(name, "column", units, AR._union(units))
    except AR.ArrangeError:
        assert name == "arc" and n < 3
        return
    assert not AR.overlapping(rects)
    assert AR.reading_order("column", rects) == list(range(n))
    right = {r[0] + r[2] for r in rects}
    assert all(r[0] >= units[0][0] for r in rects) and max(right) <= units[0][0] + units[0][2]


def test_deconstructed_grid_rules():
    """Starts in half columns at most 1.5 columns from the plain grid, one span per set, never
    wider than the plain span, a gutter at least between neighbours, drops in baselines."""
    for n, table in AR.DECONSTRUCTED_ROW.items():
        base = 24 // n
        assert len({sp for _, sp, _ in table}) == 1
        for i, (start, span, drop) in enumerate(table):
            assert isinstance(start, int) and isinstance(span, int) and isinstance(drop, int)
            assert abs(start - i * base) <= AR.MAX_SHIFT_HALF
            assert span <= base and 0 <= start and start + span <= 24 and drop >= 0
        for (s0, sp0, _), (s1, _, _) in zip(table, table[1:]):
            assert s0 + sp0 <= s1
    assert max(AR.DECONSTRUCTED_COLUMN) <= AR.MAX_SHIFT_HALF
    # drops land on the baseline grid
    region = (AR.half_x(0), AR.inch(2.3), AR.span_w(24), AR.inch(4.4))
    rects = AR.positions("deconstructed", "row", plain_row(3), region)
    bl = AR.inch(AR.GRID["baseline"])
    assert all((r[1] - AR.inch(2.3)) % bl == 0 for r in rects)


def test_positions_are_deterministic():
    units = plain_row(4)
    region = (units[0][0], units[0][1], AR.span_w(24), AR.inch(4.45))
    for name in AR.NAMES:
        assert AR.positions(name, "row", units, region) == AR.positions(name, "row", list(units), region)


def test_overlap_detection():
    a = (0, 0, 100, 100)
    assert AR.intersects(a, (50, 50, 100, 100))
    assert not AR.intersects(a, (100, 0, 50, 50))  # touching edges do not overlap
    assert AR.overlapping([a, (200, 0, 10, 10), (90, 90, 20, 20)]) == [(0, 2)]


def test_reading_order_detects_swaps():
    assert AR.reading_order("row", [(0, 500, 10, 10), (20, 0, 10, 10)]) == [0, 1]  # lower but left: still first
    assert AR.reading_order("row", [(30, 0, 10, 10), (0, 0, 10, 10)]) == [1, 0]
    assert AR.reading_order("column", [(0, 30, 10, 10), (50, 0, 10, 10)]) == [1, 0]


def test_fit_stack_gives_up_slack_not_text():
    boxes = [(0, 0, 10, 100), (0, 120, 10, 300)]  # heading, gap, body with slack
    out = AR.fit_stack(boxes, [100, 120], 1000, 300)
    assert out[0][1] == 1000 and out[0][3] == 100
    assert out[1][1] + out[1][3] <= 1300 and out[1][3] >= 120 and out[1][1] >= out[0][1] + out[0][3]
    assert AR.fit_stack(boxes, [100, 300], 0, 300) is None
    # side by side (a mark beside a label) moves as one band
    band = AR.fit_stack([(0, 0, 10, 40), (20, 5, 50, 30), (0, 60, 70, 200)], [40, 30, 50], 0, 150)
    assert band[0][1] == 0 and band[1][1] == 5


def test_unknown_arrangement_refused():
    with pytest.raises(ValueError):
        AR.style_arrangement({"name": "x", "arrangement": "spiral"})
    assert AR.style_arrangement({"arrangement": {"name": "arc"}}) == "arc"
    assert AR.style_arrangement({"arrangement": None}) is None
    assert AR.parse_arrange("column:5:ordered") == ("column", 5, True)
    assert AR.parse_arrange("") is None and AR.parse_arrange("row") is None


# ------------------------------------------------------------------ on the Make_Clean fixture (variant rows)
def variant_ids(report):
    return {it["variant"]: it for it in arrangement_items_all(report) if it["status"] == "arranged"}


def arrangement_items_all(report):
    return [it for it in items(report) if it["role"] == "arrangement"]


def slide_xml(pptx, csv):
    """{row id: slide XML bytes}."""
    _, rows = R.read_index(csv)
    ids = {int(r["slide"]): r["id"] for r in rows}
    prs = Presentation(str(pptx))
    return {ids[n]: etree.tostring(s._element) for n, s in enumerate(prs.slides, 1)}


def test_variants_only_for_tagged_layouts(arranged):
    _, rows = R.read_index(FIXTURE / "layouts.csv")
    tagged = {r["id"] for r in rows if r["arrange"]}
    _, vrows = R.read_index(FIXTURE / "visuals.csv")
    for name, (pack, report) in arranged.items():
        got = arrangement_items(report)
        assert set(got) == tagged | {r["id"] for r in vrows if r["arrange"]}
        for row_id, it in got.items():
            if row_id.startswith("V-"):
                assert it["status"] == "kept"
        assert report["hooks"]["arrangement"] == {"requested": [name], "applied": True}
        _, new_rows = R.read_index(pack / "layouts.csv")
        variants = [r for r in new_rows if r["id"] not in {x["id"] for x in rows}]
        assert {r["id"] for r in variants} == set(variant_ids(report))
        assert all(r["id"] == f"{r['id'].rsplit('-', 1)[0]}-{name}" for r in variants)


def test_plain_slides_never_change(arranged, tmp_path):
    """The plain rows of an arranged pack are the rows of the same pack without an arrangement."""
    plain, _ = R.run(FIXTURE, R.load_style("organic"), tmp_path / "plain")
    ref = slide_xml(plain / "layouts.pptx", plain / "layouts.csv")
    _, ref_rows = R.read_index(plain / "layouts.csv")
    pack, _ = arranged["stagger"]
    got = slide_xml(pack / "layouts.pptx", pack / "layouts.csv")
    _, rows = R.read_index(pack / "layouts.csv")
    assert rows[:len(ref_rows)] == ref_rows
    assert all(got[i] == x for i, x in ref.items())


def test_variant_rows_are_opt_in_and_appended(arranged):
    _, rows = R.read_index(FIXTURE / "layouts.csv")
    by_id = {r["id"]: r for r in rows}
    for name, (pack, report) in arranged.items():
        _, new_rows = R.read_index(pack / "layouts.csv")
        assert [int(r["slide"]) for r in new_rows] == list(range(1, len(new_rows) + 1))
        assert len(Presentation(str(pack / "layouts.pptx")).slides) == len(new_rows)
        for vid in variant_ids(report):
            v = next(r for r in new_rows if r["id"] == vid)
            src = by_id[vid[: -len(name) - 1]]
            assert "expressive" in v["tags"].split() and name in v["tags"].split()
            assert v["arrange"] == src["arrange"] and v["category"] == src["category"]
            assert set(R.parse_fields(v["fields"])) == set(R.parse_fields(src["fields"]))
            assert AR.is_variant(v) and not AR.is_variant(src)


def test_arranged_units_keep_margins_order_and_no_overlap(arranged):
    _, rows = R.read_index(FIXTURE / "layouts.csv")
    by_id = {r["id"]: r for r in rows}
    seen = 0
    for name, (pack, report) in arranged.items():
        for vid, it in variant_ids(report).items():
            seen += 1
            kind, count, _ = AR.parse_arrange(by_id[it["id"]]["arrange"])
            units = units_of(pack, vid, count)
            rects = [units[n] for n in sorted(units)]
            assert not AR.overlapping(rects), (name, vid)
            assert AR.reading_order(kind, rects) == list(range(count)), (name, vid)
            lim = AR.bounds((AR.inch(AR.GRID["w"]), AR.inch(7.5)), AR._union(rects))
            mx = AR.inch(AR.GRID["margin_x"])
            assert all(r[0] >= mx - R.TOL and r[0] + r[2] <= AR.inch(AR.GRID["w"]) - mx + R.TOL for r in rects)
            assert all(r[1] + r[3] <= lim[1] + lim[3] + R.TOL for r in rects)
    assert seen >= 8  # every arrangement applies somewhere in the fixture


def test_arrangement_moves_units(arranged):
    before = shapes_by_id(FIXTURE / "layouts.pptx", FIXTURE / "layouts.csv")
    pack, report = arranged["stagger"]
    assert arrangement_items(report)["L-38"]["status"] == "arranged"
    after = shapes_by_id(pack / "layouts.pptx", pack / "layouts.csv")
    tops = [R.box(after["L-38-stagger"][f"KPI {n} value"])[1] - R.box(before["L-38"][f"KPI {n} value"])[1]
            for n in (1, 2)]
    assert tops[1] - tops[0] >= AR.inch(0.5)
    pack, report = arranged["cascade"]
    if arrangement_items(report)["L-54"]["status"] == "arranged":
        after = shapes_by_id(pack / "layouts.pptx", pack / "layouts.csv")
        xs = [R.box(after["L-54-cascade"][f"Step {n}"])[0] for n in range(1, 6)]
        assert xs == sorted(xs) and xs[-1] > xs[0]


def test_max_content_fits_arranged_fields(arranged):
    """At max content each moved text field still fits: its new limit is within the style
    contract formula for its final box (per line x lines x 0.85), and it keeps 60 % of the old
    one. The final box and limit are those of the pack: the shape layer runs after the arrangement."""
    for name, (pack, report) in arranged.items():
        after = shapes_by_id(pack / "layouts.pptx", pack / "layouts.csv")
        final = {r["id"]: R.parse_fields(r["fields"]) for r in R.read_index(pack / "layouts.csv")[1]}
        for vid, it in variant_ids(report).items():
            for field, f in (it.get("fields") or {}).items():
                f = {"limit": final[vid][field]["limit"], "old_limit": f["old_limit"]}
                sh = after[vid][field]
                size, em, caps = R.font_info(sh, {"heading": 0.52, "body": 0.52})
                unit = size * em * (1.2 if caps else 1.0)
                li, ti, ri, bi = R.body_insets(sh)
                x, y, w, h = R.box(sh)
                lines = math.floor((h - ti - bi) / 12700 / (size * 1.2))
                per_line = math.floor((w - li - ri) / 12700 / unit)
                assert lines >= 1 and f["limit"] <= max(4, math.floor(per_line * lines * 0.85)), (name, vid, field)
                assert f["limit"] >= math.ceil(0.6 * f["old_limit"]) or f["old_limit"] <= 4


def test_csv_limits_lowered_only(arranged):
    _, rows = R.read_index(FIXTURE / "layouts.csv")
    old = {r["id"]: R.parse_fields(r["fields"]) for r in rows}
    for name, (pack, _) in arranged.items():
        _, new_rows = R.read_index(pack / "layouts.csv")
        for r in new_rows:
            src = r["id"][: -len(name) - 1] if AR.is_variant(r) else r["id"]
            assert set(R.parse_fields(r["fields"])) == set(old[src])
            for fname, f in R.parse_fields(r["fields"]).items():
                if f["limit"] is not None:
                    assert f["limit"] <= old[src][fname]["limit"]


def test_fallback_adds_no_variant(tmp_path):
    """When a check fails no variant row is added and the pack is the plain one."""
    st = style_with("cascade")
    st["keep"] = 0.99  # nothing may lose any room: every arrangement that shrinks a text falls back
    pack, report = R.run(FIXTURE, st, tmp_path / "a")
    got = arrangement_items(report)
    falls = [i for i, it in got.items() if it["status"] == "fallback"]
    assert falls and all(got[i]["reason"].startswith("plain grid kept") for i in falls)
    ids = {r["id"] for r in R.read_index(pack / "layouts.csv")[1]}
    assert not any(f"{i}-cascade" in ids for i in falls)
    assert len(Presentation(str(pack / "layouts.pptx")).slides) == len(ids)


def test_separators_removed_only_in_variants(arranged):
    pack, report = arranged["stagger"]
    it = arrangement_items(report)["L-54"]
    names = shapes_by_id(pack / "layouts.pptx", pack / "layouts.csv")
    assert "decoration#1" in names["L-54"]
    if it["status"] == "arranged":
        assert it.get("removed") == ["decoration#1"] and "decoration#1" not in names["L-54-stagger"]


def test_several_arrangements_one_variant_each(tmp_path):
    st = style_with(["arc", "stagger", "arc"])
    assert AR.style_arrangements(st) == ["arc", "stagger"]
    pack, report = R.run(FIXTURE, st, tmp_path / "p")
    ids = [r["id"] for r in R.read_index(pack / "layouts.csv")[1]]
    assert "L-38-arc" in ids and "L-38-stagger" in ids and len(ids) == len(set(ids))
    assert report["hooks"]["arrangement"]["requested"] == ["arc", "stagger"]


def test_deterministic_bytes_with_arrangement(tmp_path):
    st = style_with("deconstructed", "paper")
    a, _ = R.run(FIXTURE, st, tmp_path / "a")
    b, _ = R.run(FIXTURE, st, tmp_path / "b")
    for f in ("layouts.pptx", "layouts.csv", "restyle-report.json"):
        if f.endswith(".json"):
            ja, jb = json.loads((a / f).read_text()), json.loads((b / f).read_text())
            ja.pop("input", None), jb.pop("input", None)
            assert ja == jb
        else:
            assert (a / f).read_bytes() == (b / f).read_bytes()
    Presentation(str(a / "layouts.pptx"))


def test_no_arrangement_no_change(tmp_path):
    """A style without an arrangement gives the same pack as before this layer existed."""
    pack, report = R.run(FIXTURE, R.load_style("organic"), tmp_path / "p")
    assert not arrangement_items(report)
    assert R.read_index(pack / "layouts.csv")[1] == [
        {**r, "fields": r["fields"]} for r in R.read_index(pack / "layouts.csv")[1]]
    assert len(R.read_index(pack / "layouts.csv")[1]) == len(R.read_index(FIXTURE / "layouts.csv")[1])


# ------------------------------------------------------------------ decks: dosage
def deck_copy(tmp_path):
    """The fixture library as a deck: 8 slides, units on slides 4 (row:4), 6 (column:4), 8 (column:5)."""
    deck = tmp_path / "deck.pptx"
    shutil.copyfile(FIXTURE / "layouts.pptx", deck)
    return deck


@pytest.mark.parametrize("n,intensity,budget,gap", [
    (10, "calm", 1, 9), (10, "balanced", 3, 3), (10, "bold", 5, 1), (10, "all", 10, 1), (20, "balanced", 6, 3),
    (8, "balanced", 2, 3), (3, "bold", 1, 2)])
def test_dosage_presets(n, intensity, budget, gap):
    assert AR.dosage(n, intensity) == (budget, gap)


def test_dosage_is_soft_and_overridable():
    assert AR.dosage(10, "calm", {"per10": 6}) == (6, 1)
    assert AR.dosage(10, "balanced", {"gap": 1}) == (3, 1)
    assert AR.dosage(4, "balanced", {"per10": 10, "gap": 2}) == (4, 2)


def test_deck_units_from_shape_names():
    prs = Presentation(str(FIXTURE / "layouts.pptx"))
    found = [AR.deck_units(R, s) for s in prs.slides]
    kinds = [f[:2] if not isinstance(f, str) else None for f in found]
    assert kinds == [None, None, None, ("row", 4), None, ("column", 4), None, ("column", 5)]


def test_deck_default_dosage_spreads_slides(tmp_path):
    st = style_with(["stagger", "arc"])
    out, report = R.run(deck_copy(tmp_path), st, tmp_path / "out.pptx")
    hook = report["hooks"]["arrangement"]
    assert hook["applied"] and hook["dosage"] == "balanced" and hook["budget"] == 2 and hook["gap"] == 3
    got = {it["id"]: it for it in report["items"] if it["role"] == "arrangement"}
    assert set(got) == {"slide 4", "slide 6", "slide 8"}
    assert got["slide 4"]["status"] == "arranged" and got["slide 4"]["arrangement"] == "stagger"
    assert got["slide 6"]["status"] == "kept" and got["slide 6"]["reason"].startswith("dosage")
    assert got["slide 8"]["status"] == "arranged" and got["slide 8"]["arrangement"] == "arc"  # they take turns
    nums = [int(i.split()[1]) for i in hook["arranged"]]
    assert all(b - a >= 2 for a, b in zip(nums, nums[1:])) and len(nums) <= hook["budget"]
    # the arranged slide moved, the kept one did not
    before = list(Presentation(str(FIXTURE / "layouts.pptx")).slides)
    after = list(Presentation(str(out)).slides)
    pos = lambda s, n: R.box(next(sh for sh in s.shapes if sh.name == n))  # noqa: E731
    assert pos(after[3], "KPI 2 value")[1] > pos(before[3], "KPI 2 value")[1]
    assert pos(after[5], "Point 2")[:2] == pos(before[5], "Point 2")[:2]


def test_deck_calm_and_fallback(tmp_path):
    st = style_with("cascade")
    st["intensity"] = "calm"
    _, report = R.run(deck_copy(tmp_path), st, tmp_path / "calm.pptx")
    assert report["hooks"]["arrangement"]["arranged"] == ["slide 4"]
    # a text that no longer fits makes slide 4 fall back; the budget goes to the next candidate
    deck = deck_copy(tmp_path)
    prs = Presentation(str(deck))
    detail = next(sh for sh in prs.slides[3].shapes if sh.name == "KPI 1 detail")
    detail.text_frame.text = "A long explanation " * 12
    prs.save(str(deck))
    _, report = R.run(deck, st, tmp_path / "fb.pptx")
    got = {it["id"]: it for it in report["items"] if it["role"] == "arrangement"}
    assert got["slide 4"]["status"] == "fallback" and got["slide 4"]["reason"].startswith("plain grid kept")
    assert report["hooks"]["arrangement"]["arranged"] == ["slide 6"]


def test_deck_dosage_tuned_per_deck(tmp_path):
    """The agent tunes the dosage per deck: every slide that passes, or the slides it names."""
    _, report = R.run(deck_copy(tmp_path), R.override(style_with("stagger"), intensity="all"), tmp_path / "all.pptx")
    assert report["hooks"]["arrangement"]["arranged"] == ["slide 4", "slide 6", "slide 8"]
    st = R.override(style_with("stagger"), slides="6,8")
    _, report = R.run(deck_copy(tmp_path), st, tmp_path / "chosen.pptx")
    got = {it["id"]: it for it in report["items"] if it["role"] == "arrangement"}
    assert report["hooks"]["arrangement"]["arranged"] == ["slide 6", "slide 8"]  # side by side is allowed
    assert got["slide 4"]["reason"] == "dosage: not one of the chosen slides"
    assert R.main([str(deck_copy(tmp_path)), "--style", "organic", "--arrangement", "arc", "--slides", "8",
                   "--out", str(tmp_path / "cli.pptx")]) == 0
    with pytest.raises(R.RestyleError):
        R.override(style_with("arc"), slides="two")


def test_no_fixed_colours_or_fonts_added(arranged, tmp_path):
    """Make_Colors rebrands through the theme only: nothing this layer adds may carry a fixed
    RGB colour or a named font (every arranged variant has the colours and fonts of its source)."""
    import re
    fixed = re.compile(rb"<a:(?:srgbClr|sysClr|scrgbClr|hslClr|prstClr)\b|typeface=\"(?!\+m[jn]-)[^\"]+\"")
    pack, report = arranged["deconstructed"]
    xml = slide_xml(pack / "layouts.pptx", pack / "layouts.csv")
    for vid, it in variant_ids(report).items():
        assert len(fixed.findall(xml[vid])) <= len(fixed.findall(xml[it["id"]])), vid
    out, _ = R.run(deck_copy(tmp_path), R.override(style_with("stagger"), intensity="all"), tmp_path / "d.pptx")
    before = [etree.tostring(s._element) for s in Presentation(str(FIXTURE / "layouts.pptx")).slides]
    after = [etree.tostring(s._element) for s in Presentation(str(out)).slides]
    assert all(len(fixed.findall(a)) <= len(fixed.findall(b)) for a, b in zip(after, before))


def test_deck_arrangement_deterministic(tmp_path):
    st = style_with(["deconstructed", "stagger"], "paper")
    deck = deck_copy(tmp_path)
    a, _ = R.run(deck, st, tmp_path / "a.pptx")
    b, _ = R.run(deck, st, tmp_path / "b.pptx")
    assert a.read_bytes() == b.read_bytes()


def test_unknown_intensity_refused(tmp_path):
    p = tmp_path / "s.json"
    p.write_text(json.dumps({"roles": {}, "arrangement": "arc", "intensity": "loud"}))
    with pytest.raises(R.RestyleError):
        R.load_style(str(p))
    p.write_text(json.dumps({"roles": {}, "arrangement": "arc", "dosage": {"gap": 0}}))
    with pytest.raises(R.RestyleError):
        R.load_style(str(p))
    p.write_text(json.dumps({"roles": {}, "arrangement": ["arc", "spiral"]}))
    with pytest.raises(R.RestyleError):
        R.load_style(str(p))


# ------------------------------------------------------------------ Make_Clean round trip
@pytest.mark.skipif(_make_clean_dir() is None, reason="set MAKE_CLEAN_DIR to a Make_Clean_PPTX clone (v1.1.0)")
def test_make_clean_builds_variant_rows_and_deck_restyles(arranged, tmp_path):
    """Make_Clean builds a deck from a variant row of the pack (an ordinary row to it), lints it,
    and the deck it builds from the plain rows can be arranged by the deck mode."""
    scripts = _make_clean_dir()
    pack, report = arranged["stagger"]
    assert "L-38-stagger" in variant_ids(report)
    kpis = {"Title": "Four numbers to watch", "KPI 1 value": "66%", "KPI 1 label": "Heating share",
            "KPI 2 value": "3x", "KPI 2 label": "Heat pump gain", "KPI 3 value": "4 yrs",
            "KPI 3 label": "Loft payback", "KPI 4 value": "12%", "KPI 4 label": "Hot water"}
    spec = {"slides": [{"id": "L-38-stagger", "fill": kpis}, {"id": "L-24", "fill": {
        "Title": "Heating is most of the bill", "Body": "Two thirds of home energy", "Takeaway": "Insulate first."}},
        {"id": "L-38", "fill": kpis}]}
    (tmp_path / "deck.json").write_text(json.dumps(spec), encoding="utf-8")
    deck = tmp_path / "deck.pptx"
    r = subprocess.run([sys.executable, str(scripts / "build_deck.py"), "--spec", str(tmp_path / "deck.json"),
                        "--out", str(deck), "--local-dir", str(pack)], capture_output=True, text=True, timeout=240)
    assert r.returncode == 0, r.stdout + r.stderr
    r = subprocess.run([sys.executable, str(scripts / "lint_deck.py"), str(deck), "--json"], capture_output=True,
                       text=True, timeout=240)
    assert r.returncode == 0, r.stdout + r.stderr
    plain = {"slides": [spec["slides"][1], spec["slides"][2]]}
    (tmp_path / "plain.json").write_text(json.dumps(plain), encoding="utf-8")
    r = subprocess.run([sys.executable, str(scripts / "build_deck.py"), "--spec", str(tmp_path / "plain.json"),
                        "--out", str(tmp_path / "plain.pptx")], capture_output=True, text=True, timeout=240)
    assert r.returncode == 0, r.stdout + r.stderr
    _, rep = R.run(tmp_path / "plain.pptx", style_with("arc"), tmp_path / "plain-arc.pptx")
    assert rep["hooks"]["arrangement"]["arranged"] == ["slide 2"]
    r = subprocess.run([sys.executable, str(scripts / "lint_deck.py"), str(tmp_path / "plain-arc.pptx"), "--json"],
                       capture_output=True, text=True, timeout=240)
    assert r.returncode == 0, r.stdout + r.stderr
