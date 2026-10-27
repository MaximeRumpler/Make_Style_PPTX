"""Restyler: tags, geometry, safe text boxes, fallback, logo, determinism, round trips."""
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

import _shapes as S
import restyle as R

A = "{%s}" % S.A_NS
FIXTURE = Path(__file__).resolve().parent / "fixtures" / "make-clean-1.1.0"
STYLES = ("organic", "paper")


@pytest.fixture(scope="module")
def packs(tmp_path_factory):
    out = {}
    for name in STYLES:
        dst = tmp_path_factory.mktemp(name) / "pack"
        out[name] = R.run(FIXTURE, R.load_style(name), dst)
    return out


def items(report, status=None):
    for lib in report["libraries"].values():
        for it in lib["items"]:
            if status is None or it["status"] == status:
                yield it


def shapes_by_id(pptx: Path, csv: Path):
    """{row id: {shape name: shape}} over every collection of every slide."""
    _, rows = R.read_index(csv)
    ids = {int(r["slide"]): r["id"] for r in rows}
    prs = Presentation(str(pptx))
    out = {}
    for n, slide in enumerate(prs.slides, 1):
        named = {}
        for coll in R.collections(slide.shapes):
            for sh in coll:
                named[sh.name] = sh
        out[ids[n]] = named
    return out


def text_rect(sh):
    """The text rectangle of a custGeom shape, from its txL/txT/txR/txB guides."""
    gd = {g.get("name"): int(g.get("fmla").split()[2]) for g in sh._element.iter(A + "gd")}
    x, y, w, h = R.box(sh)
    return (x + w * gd["txL"] // 100000, y + h * gd["txT"] // 100000,
            w * (gd["txR"] - gd["txL"]) // 100000, h * (gd["txB"] - gd["txT"]) // 100000)


def test_tags_found_on_make_clean_library(packs):
    _, report = packs["organic"]
    roles = {it["role"] for it in items(report)}
    assert {"card", "callout", "keyfigure", "label", "panel", "logo", "decoration"} <= roles
    ids = {it["id"] for it in items(report, "restyled")}
    assert {"L-24", "L-38", "L-54", "V-PRO-04"} <= ids


def test_geometry_replaced_and_fallback_kept(packs):
    for name in STYLES:
        out, report = packs[name]
        shapes = {lib: shapes_by_id(out / f"{lib}.pptx", out / f"{lib}.csv") for lib in ("layouts", "visuals")}
        for lib, data in report["libraries"].items():
            for it in data["items"]:
                sh = shapes[lib][it["id"]][it["shape"]]
                cust = sh._element.spPr.find(A + "custGeom")
                if it["status"] == "restyled":
                    assert cust is not None, (name, it)
                else:
                    assert cust is None, (name, it)


def test_fill_and_theme_colours_kept(packs):
    out, report = packs["paper"]
    before = shapes_by_id(FIXTURE / "layouts.pptx", FIXTURE / "layouts.csv")
    after = shapes_by_id(out / "layouts.pptx", out / "layouts.csv")
    for it in report["libraries"]["layouts"]["items"]:
        if it["status"] != "restyled":
            continue
        f0 = before[it["id"]][it["shape"]]._element.spPr.find(A + "solidFill")
        f1 = after[it["id"]][it["shape"]]._element.spPr.find(A + "solidFill")
        assert etree.tostring(f0) == etree.tostring(f1)
        assert b"srgbClr" not in etree.tostring(after[it["id"]][it["shape"]]._element)


def test_text_boxes_inside_safe_area(packs):
    checked = 0
    for name in STYLES:
        out, report = packs[name]
        for lib, data in report["libraries"].items():
            shapes = shapes_by_id(out / f"{lib}.pptx", out / f"{lib}.csv")
            for it in data["items"]:
                if it["status"] != "restyled":
                    continue
                cont = shapes[it["id"]][it["shape"]]
                safe = text_rect(cont)
                for field in it["fields"]:
                    if field == it["shape"]:
                        continue  # the container's own text uses the custGeom text rectangle
                    # 100 EMU: the guides store the text box as fractions of 1/100000
                    assert R.inside(R.box(shapes[it["id"]][field]), safe, 100), (name, it["id"], field)
                    checked += 1
    assert checked > 20


def test_limits_recomputed_ids_and_fields_unchanged(packs):
    for name in STYLES:
        out, report = packs[name]
        for lib in ("layouts", "visuals"):
            h0, r0 = R.read_index(FIXTURE / f"{lib}.csv")
            h1, r1 = R.read_index(out / f"{lib}.csv")
            assert h0 == h1 and [r["id"] for r in r0] == [r["id"] for r in r1]
            for a, b in zip(r0, r1):
                fa, fb = R.parse_fields(a["fields"]), R.parse_fields(b["fields"])
                assert list(fa) == list(fb)
                for k in fa:
                    if fa[k]["limit"] is not None:
                        assert fb[k]["limit"] <= fa[k]["limit"]
        lowered = [(it["id"], f) for it in items(report, "restyled") for f, v in it["fields"].items()
                   if v["limit"] < v["old_limit"]]
        assert lowered, name
        for it in items(report, "restyled"):
            for v in it["fields"].values():
                assert v["words"] >= 1 and v["limit"] >= 4


def test_index_round_trip_is_byte_identical(tmp_path):
    for lib in ("layouts", "visuals"):
        header, rows = R.read_index(FIXTURE / f"{lib}.csv")
        R.write_index(tmp_path / "x.csv", header, rows)
        assert (tmp_path / "x.csv").read_bytes() == (FIXTURE / f"{lib}.csv").read_bytes()


def test_fallback_reported(tmp_path):
    style = R.load_style("organic")
    # drops hold three words at most and need 2 in of height at 14 pt: key figures at 40 pt cannot
    style["roles"] = {"keyfigure": {"shapes": [{"kind": "drop"}]}}
    out, report = R.run(FIXTURE, style, tmp_path / "pack")
    fb = [it for it in items(report, "fallback") if it["role"] == "keyfigure"]
    assert fb and all(it["tried"][0]["kind"] == "drop" and it["tried"][0]["reason"] for it in fb)
    on_disk = json.loads((out / "restyle-report.json").read_text(encoding="utf-8"))
    assert on_disk["summary"]["by_role"]["keyfigure"]["fallback"] == len(fb)
    # the fixture's real fallback: L-50's main tile cannot keep its label limit in any shape
    _, organic = R.run(FIXTURE, R.load_style("organic"), tmp_path / "organic")
    assert any(it["id"] == "L-50" and it["role"] == "keyfigure" for it in items(organic, "fallback"))


def test_out_of_scope_untouched(packs):
    _, report = packs["organic"]
    by_id = {}
    for it in items(report):
        by_id.setdefault(it["id"], set()).add(it["status"])
    assert by_id.get("L-41", set()) <= {"kept", "out-of-scope"}  # table
    assert by_id.get("V-CMP-01", set()) <= {"kept", "out-of-scope"}  # chart
    out, _ = packs["organic"]
    a = Presentation(str(FIXTURE / "layouts.pptx")).slides
    b = Presentation(str(out / "layouts.pptx")).slides
    _, rows = R.read_index(FIXTURE / "layouts.csv")
    n = next(int(r["slide"]) for r in rows if r["id"] == "L-41")
    assert etree.tostring(a[n - 1]._element) == etree.tostring(b[n - 1]._element)


def test_logo_never_restyled(packs, tmp_path):
    for name in STYLES:
        out, report = packs[name]
        logos = [it for it in items(report) if it["role"] == "logo"]
        assert logos and all(it["status"] == "kept" for it in logos)
        before = shapes_by_id(FIXTURE / "layouts.pptx", FIXTURE / "layouts.csv")
        after = shapes_by_id(out / "layouts.pptx", out / "layouts.csv")
        for it in logos:
            assert etree.tostring(before[it["id"]]["Logo"]._element) == etree.tostring(after[it["id"]]["Logo"]._element)
    style = json.loads((R.STYLES_DIR / "organic.json").read_text(encoding="utf-8"))
    style["roles"]["logo"] = {"shapes": [{"kind": "circle"}]}
    p = tmp_path / "logo.json"
    p.write_text(json.dumps(style), encoding="utf-8")
    with pytest.raises(R.RestyleError):
        R.load_style(str(p))


def test_deterministic_bytes(tmp_path):
    for name in STYLES:
        a, _ = R.run(FIXTURE, R.load_style(name), tmp_path / f"{name}-a")
        b, _ = R.run(FIXTURE, R.load_style(name), tmp_path / f"{name}-b")
        files = sorted(p.name for p in a.iterdir())
        assert files == sorted(p.name for p in b.iterdir())
        for f in files:
            assert (a / f).read_bytes() == (b / f).read_bytes(), (name, f)


def test_seeds_differ_per_container(packs):
    _, report = packs["paper"]
    seeds = [it["params"]["seed"] for it in items(report, "restyled") if "seed" in it.get("params", {})]
    assert len(set(seeds)) == len(seeds) > 3


def test_pack_reopens(packs):
    for name in STYLES:
        out, _ = packs[name]
        assert {"layouts.pptx", "layouts.csv", "visuals.pptx", "visuals.csv", "theme-contract.json",
                "restyle.json", "restyle-report.json"} <= {p.name for p in out.iterdir()}
        for lib in ("layouts", "visuals"):
            prs = Presentation(str(out / f"{lib}.pptx"))
            _, rows = R.read_index(out / f"{lib}.csv")
            assert len(prs.slides) == len(rows)


def test_never_overwrites(tmp_path):
    style = R.load_style("organic")
    with pytest.raises(R.RestyleError):
        R.run(FIXTURE, style, FIXTURE)
    with pytest.raises(R.RestyleError):
        R.run(FIXTURE, style, FIXTURE / "inside")
    with pytest.raises(R.RestyleError):
        R.run(FIXTURE, style, FIXTURE.parent)
    out, _ = R.run(FIXTURE, style, tmp_path / "pack")
    with pytest.raises(R.RestyleError):
        R.run(FIXTURE, style, out)
    R.run(FIXTURE, style, out, force=True)


def test_deck_mode_and_hooks(tmp_path):
    deck = tmp_path / "deck.pptx"
    shutil.copyfile(FIXTURE / "layouts.pptx", deck)
    style = R.load_style("paper")
    style["arrangement"] = "stagger"
    out, report = R.run(deck, style)
    assert out.name == "deck-paper.pptx" and out.is_file()
    assert report["kind"] == "deck" and report["summary"]["total"].get("restyled")
    assert report["hooks"]["arrangement"]["applied"] is False
    Presentation(str(out))
    assert R.main([str(deck), "--style", "paper", "--force"]) == 0
    assert (tmp_path / "deck-paper.report.json").is_file()


def test_unknown_style_and_shape(tmp_path):
    with pytest.raises(R.RestyleError):
        R.load_style("no-such-style")
    p = tmp_path / "bad.json"
    p.write_text(json.dumps({"roles": {"card": {"shapes": [{"kind": "star"}]}}}))
    with pytest.raises(R.RestyleError):
        R.load_style(str(p))


def test_reference_styles_valid():
    for name in STYLES:
        st = R.load_style(name)
        assert st["name"] == name and st["roles"]
        assert "logo" not in st["roles"] and "frame" not in st["roles"]


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_libreoffice_renders_pack(packs, tmp_path):
    office = shutil.which("soffice") or shutil.which("libreoffice")
    out, _ = packs["paper"]
    r = subprocess.run([office, "--headless", f"-env:UserInstallation=file://{tmp_path}/profile", "--convert-to",
                        "pdf", "--outdir", str(tmp_path), str(out / "layouts.pptx")], capture_output=True, timeout=240)
    pdf = tmp_path / "layouts.pdf"
    assert r.returncode == 0 and pdf.is_file()
    pages = len(re.findall(rb"/Type\s*/Page[^s]", pdf.read_bytes()))
    assert pages == len(Presentation(str(out / "layouts.pptx")).slides)


def _make_clean_dir():
    d = os.environ.get("MAKE_CLEAN_DIR")
    if d and (Path(d) / "skills" / "make-clean-pptx" / "scripts" / "build_deck.py").is_file():
        return Path(d) / "skills" / "make-clean-pptx" / "scripts"
    return None


@pytest.mark.skipif(_make_clean_dir() is None, reason="set MAKE_CLEAN_DIR to a Make_Clean_PPTX clone (v1.1.0)")
def test_make_clean_round_trip(packs, tmp_path):
    scripts = _make_clean_dir()
    out, _ = packs["organic"]
    spec = {"slides": [
        {"id": "L-24", "fill": {"Title": "Heating is most of the bill",
                                "Body": "Two thirds of home energy\nHot water next",
                                "Takeaway": "Insulate before you change the boiler."}},
        {"id": "L-38", "fill": {"Title": "Four numbers to watch", "KPI 1 value": "66%", "KPI 1 label": "Heating share",
                                "KPI 2 value": "3x", "KPI 2 label": "Heat pump gain", "KPI 3 value": "4 yrs",
                                "KPI 3 label": "Loft payback", "KPI 4 value": "12%", "KPI 4 label": "Hot water"}},
        {"id": "L-47", "fill": {"Title": "In summary", "Point 1": "Heating first", "Point 2": "Then hot water",
                                "Decision": "Approve the insulation budget."}},
    ]}
    (tmp_path / "deck.json").write_text(json.dumps(spec), encoding="utf-8")
    deck = tmp_path / "deck.pptx"
    r = subprocess.run([sys.executable, str(scripts / "build_deck.py"), "--spec", str(tmp_path / "deck.json"),
                        "--out", str(deck), "--local-dir", str(out)], capture_output=True, text=True, timeout=240)
    assert r.returncode == 0, r.stdout + r.stderr
    assert len(Presentation(str(deck)).slides) == 3
    r = subprocess.run([sys.executable, str(scripts / "lint_deck.py"), str(deck), "--json"], capture_output=True,
                       text=True, timeout=240)
    assert r.returncode == 0, r.stdout + r.stderr
