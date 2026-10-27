"""Shapes library: valid geometry, determinism, safe text box, aspect bounds, rendering."""
import io
import math
import shutil

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

import _shapes as S
import contact_sheet as CS

A = "{%s}" % S.A_NS
IN = S.EMU_PER_INCH

# every kind with its defaults, plus every variant drawn on the contact sheet
DEFAULT_ASPECT = {"rectangle": 1.6, "rounded": 1.5, "pill": 3.0, "ellipse": 1.4, "circle": 1.0, "ticket": 2.0,
                  "bubble": 1.5, "torn": 1.3, "handdrawn": 1.5}
CASES = [(k, {"svg": CS.SHEET[-1][2]["svg"]} if k == "icon" else {}, DEFAULT_ASPECT.get(k)) for k in S.kinds()]
CASES += [(kind, params, aspect) for _, kind, params, aspect in CS.SHEET]
IDS = [f"{k}-{i}" for i, (k, _, _) in enumerate(CASES)]


def box_for(kind, params, aspect, height=1.5 * IN):
    a = aspect if aspect is not None else S.natural_aspect(kind, **params)
    return (0, 0, int(height * a), int(height))


def point_in(polys, x, y):
    """Even-odd point in polygon test over several rings."""
    inside = False
    for poly in polys:
        n = len(poly)
        for i in range(n):
            (x1, y1), (x2, y2) = poly[i], poly[(i + 1) % n]
            if (y1 > y) != (y2 > y) and x < x1 + (y - y1) * (x2 - x1) / (y2 - y1):
                inside = not inside
    return inside


@pytest.mark.parametrize("kind,params,aspect", CASES, ids=IDS)
def test_valid_custgeom(kind, params, aspect):
    box = box_for(kind, params, aspect)
    geom, ins = S.make_shape(kind, box, **params)
    assert geom.tag == A + "custGeom"
    names = [c.tag[len(A):] for c in geom]
    assert names == ["avLst", "gdLst", "ahLst", "cxnLst", "rect", "pathLst"]  # schema order
    paths = geom.findall(A + "pathLst/" + A + "path")
    assert paths
    for p in paths:
        assert int(p.get("w")) == box[2] and int(p.get("h")) == box[3]
        assert p[0].tag == A + "moveTo"
        assert any(c.tag == A + "cubicBezTo" for c in p) or kind in ("rectangle", "torn")
        for pt in p.iter(A + "pt"):
            assert pt.get("x").lstrip("-").isdigit() and pt.get("y").lstrip("-").isdigit()
    gds = {g.get("name"): g.get("fmla") for g in geom.iter(A + "gd")}
    assert set(gds) == {"txL", "txT", "txR", "txB"}
    assert all(f.startswith("*/ ") for f in gds.values())
    assert ins.left + ins.right < box[2] and ins.top + ins.bottom < box[3]


@pytest.mark.parametrize("kind,params,aspect", CASES, ids=IDS)
def test_safe_box_inside_geometry(kind, params, aspect):
    box = box_for(kind, params, aspect)
    w, h = box[2], box[3]
    geom = S.outline(kind, w, h, **params)
    _, ins = S.make_shape(kind, box, **params)
    polys = [S.flatten(sub, 64) for p in geom if p.fill for sub in p.subs]
    x0, y0, x1, y1 = ins.left, ins.top, w - ins.right, h - ins.bottom
    n = 40
    for i in range(n + 1):
        t = i / n
        for x, y in ((x0 + t * (x1 - x0), y0), (x0 + t * (x1 - x0), y1), (x0, y0 + t * (y1 - y0)),
                     (x1, y0 + t * (y1 - y0))):
            assert point_in(polys, x, y), f"{kind}: text box edge point ({x:.0f}, {y:.0f}) outside the shape"
    for poly in polys:  # no part of the outline pokes into the text box
        for x, y in poly:
            assert not (x0 < x < x1 and y0 < y < y1), f"{kind}: outline point inside the text box"
    # and the box is worth having: a real share of the shape
    assert (x1 - x0) * (y1 - y0) > 0.12 * w * h


@pytest.mark.parametrize("kind,params,aspect", CASES, ids=IDS)
def test_deterministic(kind, params, aspect):
    box = box_for(kind, params, aspect)
    a = etree.tostring(S.make_shape(kind, box, **params).geom)
    b = etree.tostring(S.make_shape(kind, box, **params).geom)
    assert a == b
    assert S.make_shape(kind, box, **params).insets == S.make_shape(kind, box, **params).insets


@pytest.mark.parametrize("kind", ["blob", "torn", "handdrawn"])
def test_seed_changes_shape(kind):
    box = box_for(kind, {}, DEFAULT_ASPECT.get(kind))
    a = etree.tostring(S.make_shape(kind, box, seed=1).geom)
    b = etree.tostring(S.make_shape(kind, box, seed=2).geom)
    assert a != b


def test_pinned_bytes():
    """Same parameters, same bytes, across runs and Python versions: a change of this
    hash means the drawn shapes changed (update it on purpose, with a new contact sheet)."""
    import hashlib

    h = hashlib.sha256()
    for kind, params in [("blob", {"seed": 3}), ("torn", {"seed": 3}), ("handdrawn", {"seed": 3}),
                         ("drop", {"tilt": 12}), ("cloud", {})]:
        nat = S.natural_aspect(kind, **params) or 1.4
        res = S.make_shape(kind, (0, 0, int(1.5 * IN * nat), int(1.5 * IN)), **params)
        h.update(etree.tostring(res.geom))
        h.update(repr(tuple(res.insets)).encode())
    assert h.hexdigest() == "960f2b51078ff02511a948b9149a848bc6115e88cadcb9b190cd5dd4043b6895"


@pytest.mark.parametrize("kind,params,aspect", CASES, ids=IDS)
def test_aspect_bounds(kind, params, aspect):
    lo, hi = S.aspect_bounds(kind, **params)
    assert 0 < lo < hi
    h = IN
    assert S.fits(kind, h * math.sqrt(lo * hi), h, **params)
    assert not S.fits(kind, h * lo * 0.8, h, **params)
    assert not S.fits(kind, h * hi * 1.25, h, **params)
    with pytest.raises(S.ShapeError):
        S.make_shape(kind, (0, 0, int(h * hi * 1.25), h), **params)
    S.make_shape(kind, (0, 0, int(h * hi * 1.25), h), strict=False, **params)  # opt out


def test_drop_proportions():
    """Round body about 60-65 % of the height; the tip is the top-centre point."""
    nat = S.natural_aspect("drop")
    assert 0.6 <= nat <= 0.66
    geom = S.outline("drop", 630, 1000)
    first = geom[0].subs[0][0]
    assert first[0] == "M" and abs(first[1] - 315) < 1 and abs(first[2]) < 1
    tilted = S.natural_aspect("drop", tilt=20)
    assert tilted > nat


def test_errors():
    with pytest.raises(S.ShapeError):
        S.make_shape("star", (IN, IN))
    with pytest.raises(S.ShapeError):
        S.make_shape("bubble", (2 * IN, IN), side="middle")
    with pytest.raises(S.ShapeError):
        S.make_shape("polygon", (IN, IN), sides=2)
    with pytest.raises(S.ShapeError):
        S.make_shape("torn", (IN, IN), edges="inside")
    with pytest.raises(S.ShapeError):
        S.make_shape("rounded", (0, IN))
    with pytest.raises(S.ShapeError):
        S.make_shape("icon", (IN, IN))


def test_char_limit_matches_contract():
    # 2 x 1 in at 14 pt, char_em 0.52: 19 per line, 4 lines -> floor(19 * 4 * 0.85) = 64
    assert S.char_limit(2 * IN, IN, 14, 0.52) == 64
    # capitals widen characters by 1.2: 16 per line -> floor(16 * 4 * 0.85) = 54
    assert S.char_limit(2 * IN, IN, 14, 0.52, caps=True) == 54
    assert S.char_limit(IN // 10, IN // 10, 14) == 4  # floor of 4
    ins = S.make_shape("rounded", (0, 0, 3 * IN, 2 * IN)).insets
    w, h = 3 * IN - ins.left - ins.right, 2 * IN - ins.top - ins.bottom
    assert S.text_limit("rounded", (3 * IN, 2 * IN), 12) == S.char_limit(w, h, 12)
    assert S.text_limit("drop", (630000, 1000000), 12) < S.text_limit("rectangle", (1000000, 1000000), 12)


def test_insets_helpers():
    ins = S.Insets(10, 20, 30, 40)
    assert ins.rect(100, 200, 1000, 500) == (110, 220, 960, 440)
    assert ins.fractions(1000, 500) == (0.01, 0.04, 0.03, 0.08)


def test_icon_as_container():
    svg = CS.ICONS / "tabler" / "bookmark.svg"  # has an XML comment before <svg>
    nat = S.natural_aspect("icon", svg=str(svg))
    geom, ins = S.icon_as_container(svg, (0, 0, int(IN * nat), IN))
    assert geom.find(A + "pathLst") is not None and ins.top >= 0


def all_shapes_pptx():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    n = 0
    for i, (kind, params, aspect) in enumerate(CASES):
        x, y, w, h = box_for(kind, params, aspect, height=0.9 * IN)
        shp, _ = S.add_shape(slide, kind, (Inches(0.2) + (i % 8) * Inches(1.6), Inches(0.2) + (i // 8) * IN, w, h),
                             name=f"card#{i + 1}", **params)
        shp.text_frame.text = "Text"
        n += 1
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), n


def test_python_pptx_reopens():
    data, n = all_shapes_pptx()
    prs = Presentation(io.BytesIO(data))
    shapes = list(prs.slides[0].shapes)
    assert len(shapes) == n
    for shp in shapes:
        assert shp._element.spPr.find(A + "custGeom") is not None
        assert shp._element.spPr.find(A + "prstGeom") is None
        assert shp.text_frame.text == "Text"


@pytest.mark.skipif(not CS.soffice(), reason="LibreOffice not installed")
def test_libreoffice_renders(tmp_path):
    import subprocess

    data, _ = all_shapes_pptx()
    src = tmp_path / "shapes.pptx"
    src.write_bytes(data)
    subprocess.run([CS.soffice(), "--headless", f"-env:UserInstallation=file://{tmp_path}/profile",
                    "--convert-to", "png", "--outdir", str(tmp_path), str(src)],
                   check=True, capture_output=True, timeout=240)
    png = tmp_path / "shapes.png"
    assert png.exists() and png.stat().st_size > 5000
    from PIL import Image

    im = Image.open(png).convert("L")
    dark = sum(im.histogram()[:200])
    assert dark > 0.02 * im.width * im.height  # the shapes were drawn, not dropped


@pytest.mark.skipif(not (CS.soffice() and shutil.which("pdftoppm")), reason="LibreOffice or pdftoppm missing")
def test_contact_sheet(tmp_path):
    out = CS.build(tmp_path / "sheet.pptx")
    assert len(Presentation(str(out)).slides) == math.ceil(len(CS.SHEET) / (CS.COLS * CS.ROWS))
    png = CS.render(out, tmp_path / "sheet.png", dpi=30)
    assert png and png.exists()


def test_contact_sheet_builds_without_office(tmp_path):
    out = CS.build(tmp_path / "sheet.pptx")
    prs = Presentation(str(out))
    kinds = {shp.name.split(" ")[0] for s in prs.slides for shp in s.shapes if shp.name.split(" ")[0] in S.KINDS}
    assert kinds == set(S.KINDS)
